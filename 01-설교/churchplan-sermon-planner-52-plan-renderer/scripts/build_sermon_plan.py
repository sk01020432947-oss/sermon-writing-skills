#!/usr/bin/env python3
"""
build_sermon_plan.py — churchplan-sermon-planner-52 계획표 렌더러.

sermon-plan.json → ① sermon-plan.md (52주 계획표 + 분기 시리즈 개요 + 설교–목표 지도 + 부하 경고)
                   ② sermon-plan.xlsx (52주 표·시리즈 개요·설교–목표/부하 시트, openpyxl 없으면 .csv)
                   ③ sermon-series-overview.docx (분기별 설교 시리즈 개요, python-docx 있을 때)

렌더 전 핵심 게이트를 재강제한다(cys-insight 패턴):
  필수 블록 · 표어(chosen_theme) · 주차 중복. 위반 시 렌더 거부(종료코드 1).
본문 실재·절기 일치 등 정밀 검증은 fidelity-guard의 validate_sermon_plan.py가 맡는다.

사용:  python3 build_sermon_plan.py <sermon-plan.json> [출력디렉토리]
종료코드: 0 = 성공, 1 = 게이트 위반(렌더 거부), 2 = 파일/JSON 오류
표준 라이브러리만 사용(xlsx는 openpyxl 있을 때만, 없으면 csv).
"""
import csv
import json
import os
import sys


def _d(x):
    """truthy 비-dict(리스트·문자열 등)도 안전하게 {}로 치환 — `(x or {})` 관용구 결함 차단."""
    return x if isinstance(x, dict) else {}


def _l(x):
    """비-list(문자열·정수 등)는 빈 리스트로 — 문자열 char 순회·crash 차단."""
    return x if isinstance(x, list) else []


def _ld(x):
    """list-of-dicts만 추출 — 비-list거나 비-dict 원소는 배제(적대 입력 크래시 차단)."""
    return [v for v in x if isinstance(v, dict)] if isinstance(x, list) else []


def gate(p):
    """렌더 전 핵심 게이트 — 본질 3종 + 구조 필드 형식(검증기-렌더러 parity).

    정밀 검증은 fidelity-guard가 맡지만, truthy 비-dict·비-list 오염은 렌더 산출을
    조용히 망가뜨리므로 여기서도 [형식오류]로 잡아 렌더를 거부한다(EXIT 1).
    """
    errs = []
    for blk in ("meta", "weeks"):
        if blk not in p:
            errs.append(f"필수 블록 '{blk}' 없음")
    if not isinstance(p.get("meta"), dict) and p.get("meta") is not None:
        errs.append(f"meta가 객체(dict)가 아님: {type(p.get('meta')).__name__}")
    if not str(_d(p.get("meta")).get("chosen_theme") or "").strip():
        errs.append("meta.chosen_theme 비어 있음 — 표어 먼저 확정(#2)")
    # weeks가 존재하되 비-list(문자열·정수)면 형식오류 — char 순회·빈 산출 방지
    if "weeks" in p and p.get("weeks") is not None and not isinstance(p.get("weeks"), list):
        errs.append(f"weeks가 배열(list)이 아님: {type(p.get('weeks')).__name__}")
    seen = set()
    for w in _l(p.get("weeks")):
        if not isinstance(w, dict):
            errs.append(f"weeks 원소가 객체(dict)가 아님: {type(w).__name__}"); continue
        wn = w.get("week")
        if wn in seen:
            errs.append(f"{wn}주차 중복 등록")
        seen.add(wn)
        wd = w.get("worship_direction")
        if wd is not None and not isinstance(wd, dict):
            errs.append(f"{wn}주차 worship_direction이 객체(dict)가 아님: {type(wd).__name__}")
        pts = w.get("parallel_texts")
        if pts is not None and not isinstance(pts, list):
            errs.append(f"{wn}주차 parallel_texts가 배열(list)이 아님: {type(pts).__name__}")
    if not _l(p.get("weeks")):
        errs.append("weeks 비어 있음")
    # inputs_echo는 검증기 schema상 dict — 렌더는 미사용이나 parity 위해 형식 검사
    if "inputs_echo" in p and p.get("inputs_echo") is not None and not isinstance(p.get("inputs_echo"), dict):
        errs.append(f"inputs_echo가 객체(dict)가 아님: {type(p.get('inputs_echo')).__name__}")
    # 선택 구조 필드 형식 — 검증기와 동일 parity
    sr = p.get("series")
    if sr is not None and not isinstance(sr, list):
        errs.append(f"series가 배열(list)이 아님: {type(sr).__name__}")
    for i, s in enumerate(_l(sr)):
        if not isinstance(s, dict):
            errs.append(f"series[{i}] 원소가 객체(dict)가 아님: {type(s).__name__}")
        elif s.get("goal_link") is not None and not isinstance(s.get("goal_link"), list):
            errs.append(f"series[{i}].goal_link가 배열(list)이 아님: {type(s.get('goal_link')).__name__}")
    nc = p.get("narrative_check")
    if nc is not None and not isinstance(nc, dict):
        errs.append(f"narrative_check가 객체(dict)가 아님: {type(nc).__name__}")
    sgm = _d(nc).get("sermon_goal_map")
    if sgm is not None and not isinstance(sgm, list):
        errs.append(f"narrative_check.sermon_goal_map이 배열(list)이 아님: {type(sgm).__name__}")
    for i, m in enumerate(_l(sgm)):
        if not isinstance(m, dict):
            errs.append(f"sermon_goal_map[{i}] 원소가 객체(dict)가 아님: {type(m).__name__}")
        elif m.get("serves_axes") is not None and not isinstance(m.get("serves_axes"), list):
            errs.append(f"sermon_goal_map[{i}].serves_axes가 배열(list)이 아님: {type(m.get('serves_axes')).__name__}")
    lw = p.get("load_warnings")
    if lw is not None and not isinstance(lw, list):
        errs.append(f"load_warnings가 배열(list)이 아님: {type(lw).__name__}")
    for i, lwi in enumerate(_l(lw)):
        if not isinstance(lwi, dict):
            errs.append(f"load_warnings[{i}] 원소가 객체(dict)가 아님: {type(lwi).__name__}")
        elif lwi.get("weeks") is not None and not isinstance(lwi.get("weeks"), list):
            errs.append(f"load_warnings[{i}].weeks가 배열(list)이 아님: {type(lwi.get('weeks')).__name__}")
    return errs


def fmt_worship(wd):
    # worship_direction이 truthy 비-dict(리스트·문자열)면 .get() raw traceback — _d로 차단.
    wd = _d(wd)
    if not wd:
        return ""
    parts = []
    for k in ("theme", "mood", "style"):
        v = wd.get(k)
        # dict/list 값은 셀 오염([...]·{...}) 차단 — 스칼라 방향만 출력
        if v and not isinstance(v, (list, dict)):
            parts.append(str(v))
    return " / ".join(parts)


def week_rows(p):
    rows = []
    for w in sorted(_l(p.get("weeks")), key=lambda x: x.get("week", 0) if isinstance(x, dict) else 0):
        if not isinstance(w, dict):
            continue
        # 비-문자열 원소(정수·리스트·dict)도 str 강제 — 검증기(validate_sermon_plan.py)와 동일 방어로
        # raw traceback·openpyxl ValueError 차단(garbage-in 차단 계약).
        pts = w.get("parallel_texts") or []
        if not isinstance(pts, list):
            pts = [pts]
        # dict/list 원소는 본문 셀 오염 차단(검증기와 동일 — 스칼라 본문만 join)
        parallels = ", ".join(str(t) for t in pts if not isinstance(t, (list, dict)))
        text = str(w.get("primary_text", ""))
        if parallels:
            text = f"{text} ({parallels})"
        rows.append([
            w.get("week", ""), w.get("date", ""), w.get("observance", ""),
            w.get("series_id", ""), text, w.get("title", ""),
            w.get("core_message", ""), w.get("application", ""),
            fmt_worship(w.get("worship_direction")),
            "확정" if w.get("confirmed") else "잠정",
        ])
    return rows


WK_HEADER = ["주차", "날짜", "절기", "시리즈", "본문", "제목", "핵심 메시지", "적용", "음악 방향", "확정"]


def build_md(p, rows):
    meta = _d(p.get("meta"))
    L = []
    L.append(f"# {meta.get('year', '')} 52주 설교 계획")
    L.append("")
    L.append(f"- 표어: **{meta.get('chosen_theme', '(미확정)')}**")
    L.append(f"- 스타일: {meta.get('preaching_style', '')}  ·  "
             + ("✅ 목사 확정" if p.get("confirmed") else "🙏 목사 확정 대기")
             + (f"  ·  롤링({'1Q 주별·2~4Q 시리즈' if meta.get('rolling_mode') else '전체 확정'})"))
    L.append("")

    # 분기 시리즈 개요 — series가 비-list면 빈 순회로 안전 환원
    series = _l(p.get("series"))
    if series:
        L.append("## 분기 시리즈 개요")
        L.append("")
        for s in series:
            if not isinstance(s, dict):
                continue
            gl = ", ".join(str(x) for x in _l(s.get("goal_link")) if not isinstance(x, (list, dict)))
            L.append(f"### [{s.get('quarter','')}Q] {s.get('title','')}  ({s.get('style','')})")
            L.append(f"- 목적: {s.get('purpose','')}")
            L.append(f"- 표어 연결: {s.get('theme_link','')}")
            if gl:
                L.append(f"- 섬길 목표: {gl}")
            L.append("")

    # 52주 표
    L.append("## 52주 설교 계획표")
    L.append("")
    L.append("| " + " | ".join(WK_HEADER) + " |")
    L.append("|" + "---|" * len(WK_HEADER))
    for r in rows:
        L.append("| " + " | ".join(str(c) for c in r) + " |")
    L.append("")

    # 설교–목표 지도 (schema상 list-of-dict — 비정상 형(dict 등)은 표를 건너뛰어 raw traceback 차단)
    nc = _d(p.get("narrative_check"))
    smap = _ld(nc.get("sermon_goal_map"))
    if smap:
        L.append("## 설교–목표 지도")
        L.append("")
        L.append("| 시리즈 | 섬기는 축 | 방법 |")
        L.append("|---|---|---|")
        for m in smap:
            axes = ", ".join(str(x) for x in _l(m.get("serves_axes")) if not isinstance(x, (list, dict)))
            L.append(f"| {m.get('series','')} | {axes} | {m.get('how','')} |")
        L.append("")
    if nc.get("series_flow"):
        L.append(f"> 서사 흐름: {nc['series_flow']}")
        L.append("")

    # 부하 경고 (비-dict 원소·비-list 값 배제 — xlsx 경로와 동일 방어)
    lw = _ld(p.get("load_warnings"))
    if lw:
        L.append("## 🛌 목자의 짐 — 부하·안식 경고")
        L.append("")
        for w in lw:
            wks = ", ".join(str(x) for x in _l(w.get("weeks")))
            L.append(f"- **[{w.get('type','')}]** {w.get('text','')}  _(주차: {wks})_")
        L.append("")

    L.append("> 본문 선택·해석의 최종 권한은 목사에게 있습니다. 이 표는 초안이며, 성령의 인도와 회중 형편에 따라 수정됩니다.")
    L.append("> 음악은 방향(주제·정서·스타일)이며, 특정 곡은 회중이 아는 곡에서 찬양팀과 선정합니다.")
    return "\n".join(L) + "\n"


def build_xlsx(p, rows, out_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        csv_path = out_path.rsplit(".", 1)[0] + ".csv"
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(WK_HEADER)
            for r in rows:
                w.writerow(r)
        return csv_path

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="594B3B")
    head_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "52주 설교 계획표"
    ws.append(WK_HEADER)
    for c in ws[1]:
        c.fill = head_fill; c.font = head_font; c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    widths = [6, 11, 12, 8, 24, 22, 30, 30, 26, 7]
    for i, wd in enumerate(widths, 1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # openpyxl 셀은 list/dict를 받으면 ValueError — 복합형 셀은 str 강제(garbage-in 차단).
    def _cell(v):
        return str(v) if isinstance(v, (list, dict)) else v

    series = _l(p.get("series"))
    if any(isinstance(s, dict) for s in series):
        ws2 = wb.create_sheet("분기 시리즈 개요")
        ws2.append(["분기", "시리즈", "스타일", "목적", "표어 연결", "섬길 목표"])
        for c in ws2[1]:
            c.fill = head_fill; c.font = head_font
        for s in series:
            if not isinstance(s, dict):
                continue
            gl = ", ".join(str(x) for x in _l(s.get("goal_link")) if not isinstance(x, (list, dict)))
            ws2.append([_cell(s.get("quarter", "")), _cell(s.get("title", "")), _cell(s.get("style", "")),
                        _cell(s.get("purpose", "")), _cell(s.get("theme_link", "")), gl])
        for col, wd in zip("ABCDEF", [6, 20, 8, 36, 24, 28]):
            ws2.column_dimensions[col].width = wd

    nc = _d(p.get("narrative_check"))
    smap_x = _ld(nc.get("sermon_goal_map"))
    lw = _ld(p.get("load_warnings"))
    if smap_x or lw:
        ws3 = wb.create_sheet("설교–목표·부하")
        ws3.append(["구분", "내용1", "내용2", "내용3"])
        for c in ws3[1]:
            c.fill = head_fill; c.font = head_font
        for m in smap_x:
            axes = ", ".join(str(x) for x in _l(m.get("serves_axes")) if not isinstance(x, (list, dict)))
            ws3.append(["설교–목표", _cell(m.get("series", "")), axes, _cell(m.get("how", ""))])
        for w in lw:
            ws3.append(["부하경고", _cell(w.get("type", "")), _cell(w.get("text", "")), ", ".join(str(x) for x in _l(w.get("weeks")))])
        for col, wd in zip("ABCD", [12, 16, 24, 36]):
            ws3.column_dimensions[col].width = wd

    wb.save(out_path)
    return out_path


FONT = "나눔명조"


def build_docx(p, out_path):
    """분기별 설교 시리즈 개요 docx — 대표 SKILL.md:74·마스터:148 'docx' 약속 충족.

    python-docx 없으면 None 반환(렌더는 md/xlsx로 계속). 패턴은 plan-compiler build_package.py 준용.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        return None

    meta = _d(p.get("meta"))
    sepia = RGBColor(0x59, 0x4B, 0x3B)

    def setfont(run, size, color=None, bold=False):
        run.font.name = FONT
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        run.font.bold = bold
        rPr = run._element.get_or_add_rPr()
        rF = rPr.find(qn("w:rFonts"))
        if rF is None:
            rF = rPr.makeelement(qn("w:rFonts"), {}); rPr.append(rF)
        for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            rF.set(qn(a), FONT)

    doc = Document()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(t.add_run(f"{meta.get('year', '')} 분기별 설교 시리즈 개요"), 20, sepia, True)
    th = doc.add_paragraph(); th.alignment = WD_ALIGN_PARAGRAPH.CENTER
    setfont(th.add_run(meta.get("chosen_theme", "")), 13, sepia)

    for s in _l(p.get("series")):
        if not isinstance(s, dict):
            continue
        h = doc.add_paragraph()
        setfont(h.add_run(f"[{s.get('quarter','')}Q] {s.get('title','')}  ({s.get('style','')})"), 14, sepia, True)
        if s.get("purpose"):
            b = doc.add_paragraph(); setfont(b.add_run(f"목적: {s.get('purpose','')}"), 11)
        if s.get("theme_link"):
            b = doc.add_paragraph(); setfont(b.add_run(f"표어 연결: {s.get('theme_link','')}"), 11)
        gl = ", ".join(str(x) for x in _l(s.get("goal_link")) if not isinstance(x, (list, dict)))
        if gl:
            b = doc.add_paragraph(); setfont(b.add_run(f"섬길 목표: {gl}"), 11)

    note = doc.add_paragraph()
    setfont(note.add_run("본문 선택·해석의 최종 권한은 목사에게 있습니다. 이 개요는 초안입니다."),
            8.5, RGBColor(0xA0, 0x99, 0x99))
    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        print("사용: python3 build_sermon_plan.py <sermon-plan.json> [출력디렉토리]"); sys.exit(2)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    try:
        with open(src, encoding="utf-8") as f:
            p = json.load(f)
    except FileNotFoundError:
        print(f"[오류] 파일 없음: {src}"); sys.exit(2)
    except json.JSONDecodeError as e:
        print(f"[오류] JSON 파싱 실패: {e}"); sys.exit(2)

    if not isinstance(p, dict):
        print(f"[오류] 최상위 JSON이 객체(dict)가 아님: {type(p).__name__} — sermon-plan.json 구조 확인"); sys.exit(2)

    errs = gate(p)
    if errs:
        print("=" * 60)
        print("렌더 거부 — 게이트 위반 (fidelity-guard 먼저 통과시키라)")
        print("=" * 60)
        for e in errs:
            print(f"  ✗ {e}")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    rows = week_rows(p)
    md_path = os.path.join(out_dir, "sermon-plan.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(build_md(p, rows))
    xlsx_path = build_xlsx(p, rows, os.path.join(out_dir, "sermon-plan.xlsx"))
    docx_path = None
    if p.get("series"):
        docx_path = build_docx(p, os.path.join(out_dir, "sermon-series-overview.docx"))

    print("렌더 완료 ✓")
    print(f"  · {md_path}")
    print(f"  · {xlsx_path}" + ("  (openpyxl 없어 csv 폴백)" if xlsx_path.endswith(".csv") else ""))
    if docx_path:
        print(f"  · {docx_path}")
    elif p.get("series"):
        print("  · (python-docx 미설치 — 분기 시리즈 개요 docx 생략)")
    sys.exit(0)


if __name__ == "__main__":
    main()
