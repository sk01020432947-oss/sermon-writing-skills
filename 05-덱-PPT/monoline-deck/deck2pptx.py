# -*- coding: utf-8 -*-
"""monoline 덱(HTML) → PPTX 변환기.

사용법:  python deck2pptx.py <deck.html> [out.pptx]
원리:  Edge 헤드리스로 슬라이드마다 #N 해시 URL을 열어 1920x1080 스크린샷을 찍고,
       python-pptx로 16:9 슬라이드에 전면 배치한다. (덱의 #N 해시 내비 기능 필요)
"""
import io, pathlib, re, shutil, subprocess, sys, tempfile

BROWSERS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/usr/bin/google-chrome", "/usr/bin/chromium", "/usr/bin/chromium-browser", "/usr/bin/microsoft-edge",
]


def find_browser():
    for p in BROWSERS:
        if pathlib.Path(p).exists():
            return p
    for name in ("msedge", "chrome", "google-chrome", "chromium", "microsoft-edge"):
        p = shutil.which(name)
        if p:
            return p
    sys.exit("Edge/Chrome 실행 파일을 찾을 수 없습니다.")


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    html = pathlib.Path(sys.argv[1]).resolve()
    if not html.exists():
        sys.exit(f"파일 없음: {html}")
    out = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else html.with_suffix(".pptx")

    n_slides = len(re.findall(r'class="slide[ "]', io.open(html, encoding="utf-8").read()))
    if not n_slides:
        sys.exit("slide 섹션을 찾지 못했습니다.")
    browser = find_browser()
    url = html.as_uri()

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as td:
        for i in range(1, n_slides + 1):
            png = pathlib.Path(td) / f"s{i:03d}.png"
            subprocess.run(
                [browser, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                 "--window-size=1920,1080", "--virtual-time-budget=4000",
                 f"--screenshot={png}", f"{url}?shot#{i}"],
                check=True, capture_output=True, timeout=60)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(png), 0, 0,
                                     width=prs.slide_width, height=prs.slide_height)
            print(f"  slide {i}/{n_slides}", end="\r")
    prs.save(out)
    print(f"\n완료: {out} ({n_slides} slides)")


if __name__ == "__main__":
    main()
