# orange-gradient-deck builder: spec JSON -> orange gradient corporate pptx
# usage: python build.py spec.json
import json, math, os, sys, tempfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageOps, ImageDraw

O_DEEP = RGBColor(0xE0, 0x4A, 0x17)
O_LIGHT = RGBColor(0xF5, 0x9E, 0x42)
O_MID = RGBColor(0xEA, 0x6C, 0x2B)
INK = RGBColor(0x33, 0x30, 0x2E)
GRAY = RGBColor(0x8A, 0x86, 0x82)
LIGHT = RGBColor(0xF6, 0xF4, 0xF2)
BORDER = RGBColor(0xE8, 0xE4, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
XBOLD = "Pretendard ExtraBold"
SEMI = "Pretendard SemiBold"
BODY = "Pretendard Regular"
SW, SH = 13.333, 7.5
_TMP = tempfile.mkdtemp(prefix="ogd_")


def _fonts(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def text(slide, x, y, w, h, s, size=10, font=BODY, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.25, wrap=True, spacing=None, upper=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run()
        r.text = ln.upper() if upper else ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)
        if spacing is not None:
            r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def shape(slide, kind, x, y, w, h, fill=WHITE, line_color=None, line_w=1.0):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    return sp


def grad(sp, c1=O_DEEP, c2=O_LIGHT, angle=45.0):
    sp.fill.gradient()
    stops = sp.fill.gradient_stops
    stops[0].color.rgb = c1
    stops[1].color.rgb = c2
    try:
        sp.fill.gradient_angle = angle
    except Exception:
        pass
    return sp


def grad_rect(slide, x, y, w, h, c1=O_DEEP, c2=O_LIGHT, angle=45.0, round_=False, adj=0.12):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if round_:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    sp.line.fill.background()
    return grad(sp, c1, c2, angle)


def label_in(sp, s, size=10, color=WHITE, font=XBOLD):
    tf = sp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)


def ring(slide, cx, cy, d, color=WHITE, weight=1.5):
    sp = shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=None,
               line_color=color, line_w=weight)
    return sp


def rings(slide, cx, cy, base=2.0, n=3, color=None):
    c = color or RGBColor(0xF3, 0x8B, 0x4E)
    for i in range(n):
        ring(slide, cx, cy, base + i * 0.85, c, 1.25 + (i % 2) * 0.75)


def card(slide, x, y, w, h, round_=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sp = shape(slide, kind, x, y, w, h, fill=WHITE, line_color=BORDER, line_w=1.0)
    if round_:
        try:
            sp.adjustments[0] = 0.08
        except Exception:
            pass
    return sp


def photo(slide, path, x, y, w, h, circle_=False):
    if path and os.path.isfile(path):
        im = Image.open(path)
        im = ImageOps.exif_transpose(im).convert("RGB")
        tw, th = max(int(w * 150), 50), max(int(h * 150), 50)
        im = ImageOps.fit(im, (tw, th), Image.LANCZOS)
        if circle_:
            mask = Image.new("L", (tw, th), 0)
            ImageDraw.Draw(mask).ellipse((0, 0, tw, th), fill=255)
            im.putalpha(mask)
            out = os.path.join(_TMP, f"p{abs(hash((path, tw, th, 1)))}.png")
            im.save(out)
        else:
            out = os.path.join(_TMP, f"p{abs(hash((path, tw, th)))}.jpg")
            im.save(out, quality=90)
        slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        if circle_:
            shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill=LIGHT, line_color=BORDER, line_w=1.0)
        else:
            shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=LIGHT)


def donut(slide, cx, cy, r, pct, color=O_MID, track=None):
    """Progress ring from 12 o'clock via freeform wedge + hole."""
    shape(slide, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill=(track or RGBColor(0xF3, 0xE3, 0xD8)))
    pct = max(0.0, min(float(pct), 100.0))
    if pct > 0:
        E = 914400
        steps = max(2, int(pct * 0.72))
        pts = []
        for k in range(steps + 1):
            ang = math.radians(-90 + 360.0 * pct / 100.0 * k / steps)
            pts.append((Emu(int((cx + r * math.cos(ang)) * E)), Emu(int((cy + r * math.sin(ang)) * E))))
        fb = slide.shapes.build_freeform(Emu(int(cx * E)), Emu(int(cy * E)), scale=1.0)
        fb.add_line_segments(pts, close=True)
        sp = fb.convert_to_shape()
        sp.shadow.inherit = False
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
        sp.line.fill.background()
    inner = shape(slide, MSO_SHAPE.OVAL, cx - r * 0.62, cy - r * 0.62, r * 1.24, r * 1.24, fill=WHITE)
    label_in(inner, f"{int(pct)}%", r * 22, INK)


def header(slide, spec, s, dark=False):
    fg = WHITE if dark else INK
    text(slide, 0.5, 0.35, 4.0, 0.3, spec.get("brand", ""), 9, XBOLD, (WHITE if dark else O_MID),
         spacing=200, upper=True)
    n = s.get("_page")
    if n:
        text(slide, SW - 1.0, 0.35, 0.5, 0.3, f"{n:02d}", 9, SEMI, fg, align=PP_ALIGN.RIGHT)


# ---------------- layouts ----------------

def lay_cover(prs, slide, spec, s):
    grad_rect(slide, 0, 0, SW, SH, angle=55.0)
    rings(slide, SW / 2, SH / 2 - 0.2, 2.6, 3, RGBColor(0xF4, 0x9B, 0x5A))
    rings(slide, 11.6, 1.2, 1.6, 2, RGBColor(0xEE, 0x7A, 0x3A))
    text(slide, 1.0, 1.5, SW - 2.0, 0.4, spec.get("brand", ""), 13, XBOLD, WHITE,
         align=PP_ALIGN.CENTER, spacing=300, upper=True)
    text(slide, 1.0, 2.6, SW - 2.0, 1.5, s.get("title", "COMPANY"), 60, XBOLD, WHITE,
         align=PP_ALIGN.CENTER, spacing=150, upper=True)
    if s.get("sub"):
        text(slide, 1.0, 4.35, SW - 2.0, 0.4, s["sub"], 12, SEMI, WHITE, align=PP_ALIGN.CENTER, spacing=200)


def lay_contents(prs, slide, spec, s):
    grad_rect(slide, 0, 0, SW, SH, angle=55.0)
    rings(slide, 2.6, 3.75, 2.2, 3, RGBColor(0xF4, 0x9B, 0x5A))
    header(slide, spec, s, dark=True)
    text(slide, 1.4, 3.3, 3.5, 0.8, s.get("title", "Contents"), 26, XBOLD, WHITE)
    items = s.get("items", [])[:6]
    y = 2.4
    for it in items:
        shape(slide, MSO_SHAPE.OVAL, 7.1, y + 0.09, 0.09, 0.09, fill=WHITE)
        text(slide, 7.45, y, 5.0, 0.35, it, 13, SEMI, WHITE)
        y += 0.62


def lay_split(prs, slide, spec, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
    header(slide, spec, s)
    photo(slide, s.get("image"), 0.9, 1.5, 5.3, 4.9)
    tab = grad_rect(slide, 7.0, 1.3, 5.4, 0.85, round_=True, adj=0.5)
    label_in(tab, s.get("title", ""), 16)
    if s.get("body"):
        text(slide, 7.2, 2.6, 5.2, 3.6, s["body"], 9.5, BODY, GRAY, line=1.5)
    if s.get("tag"):
        tg = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 7.2, 6.1, 1.9, 0.5, fill=LIGHT)
        tg.adjustments[0] = 0.5
        label_in(tg, s["tag"], 9, O_MID, SEMI)


def lay_project(prs, slide, spec, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
    header(slide, spec, s)
    text(slide, 0.9, 1.6, 4.5, 0.4, s.get("kicker", ""), 12, SEMI, GRAY)
    text(slide, 0.9, 2.0, 4.8, 1.0, s.get("title", ""), 30, XBOLD, INK)
    if s.get("body"):
        text(slide, 0.9, 3.2, 4.6, 3.0, s["body"], 9.5, BODY, GRAY, line=1.5)
    gc = grad_rect(slide, 6.3, 1.3, 6.1, 2.3, round_=True, adj=0.12, angle=30.0)
    text(slide, 6.7, 1.6, 5.3, 0.4, s.get("card_head", ""), 15, XBOLD, WHITE)
    if s.get("card_body"):
        text(slide, 6.7, 2.1, 5.3, 1.3, s["card_body"], 8.5, BODY, WHITE, line=1.35)
    photo(slide, s.get("image"), 7.4, 3.3, 3.4, 3.4, circle_=True)
    rings(slide, 11.6, 5.0, 1.4, 2, RGBColor(0xF0, 0xC9, 0xAF))


def lay_teamrow(prs, slide, spec, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
    grad_rect(slide, 0, 0, SW, 1.15, angle=0.0)
    text(slide, 0.5, 0.35, 6.0, 0.4, s.get("title", "Service team"), 16, XBOLD, WHITE)
    n2 = s.get("_page")
    if n2:
        text(slide, SW - 1.0, 0.4, 0.5, 0.3, f"{n2:02d}", 9, SEMI, WHITE, align=PP_ALIGN.RIGHT)
    if s.get("body"):
        text(slide, 0.9, 1.5, 11.5, 0.7, s["body"], 9, BODY, GRAY, line=1.35)
    items = s.get("items", [])[:4]
    n = max(len(items), 1)
    cw = (SW - 1.8) / n
    for i, it in enumerate(items):
        x = 0.9 + i * cw
        card(slide, x + 0.15, 2.5, cw - 0.5, 4.0)
        photo(slide, it.get("image"), x + cw / 2 - 0.95, 2.9, 1.7, 1.7, circle_=True)
        text(slide, x + 0.2, 4.8, cw - 0.6, 0.3, it.get("name", ""), 12, XBOLD, INK, align=PP_ALIGN.CENTER)
        text(slide, x + 0.2, 5.15, cw - 0.6, 0.3, it.get("role", ""), 8.5, SEMI, O_MID, align=PP_ALIGN.CENTER)
        text(slide, x + 0.35, 5.5, cw - 0.9, 0.85, it.get("body", ""), 7.5, BODY, GRAY,
             align=PP_ALIGN.CENTER, line=1.25)


def lay_stats(prs, slide, spec, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
    header(slide, spec, s)
    items = s.get("items", [])[:2]
    y = 1.7
    for i, it in enumerate(items):
        donut(slide, 2.0, y + 1.1, 1.05, it.get("pct", 50))
        text(slide, 3.5, y + 0.35, 3.0, 0.35, it.get("head", ""), 13, XBOLD, INK)
        text(slide, 3.5, y + 0.75, 3.1, 1.4, it.get("body", ""), 8.5, BODY, GRAY, line=1.35)
        y += 2.6
    gc = grad_rect(slide, 7.0, 1.7, 5.4, 2.5, round_=True, adj=0.12, angle=30.0)
    text(slide, 7.4, 2.05, 4.6, 0.4, s.get("card_head", ""), 15, XBOLD, WHITE)
    if s.get("card_body"):
        text(slide, 7.4, 2.55, 4.6, 1.4, s["card_body"], 8.5, BODY, WHITE, line=1.35)
    photo(slide, s.get("image"), 7.9, 4.15, 2.9, 2.9, circle_=True)
    rings(slide, 11.4, 5.6, 1.3, 2, RGBColor(0xF0, 0xC9, 0xAF))


def lay_bars(prs, slide, spec, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
    header(slide, spec, s)
    text(slide, 0.9, 1.3, 5.0, 0.5, s.get("title", "Performance"), 22, XBOLD, INK)
    bars = s.get("bars", [])[:14]
    if bars:
        bx, bw_area, by, bh = 0.9, 7.2, 6.3, 3.6
        vmax = max(b.get("value", 0) for b in bars) or 1
        bw = bw_area / len(bars) * 0.55
        gap = bw_area / len(bars)
        for i, b in enumerate(bars):
            h = bh * b.get("value", 0) / vmax
            grad_rect(slide, bx + i * gap, by - h, bw, h, c1=O_LIGHT, c2=O_DEEP, angle=90.0,
                      round_=True, adj=0.4)
            if b.get("label"):
                text(slide, bx + i * gap - 0.15, by + 0.1, gap, 0.25, b["label"], 6.5, SEMI, GRAY,
                     align=PP_ALIGN.CENTER)
        ln = slide.shapes.add_connector(1, Inches(bx - 0.1), Inches(by), Inches(bx + bw_area), Inches(by))
        ln.line.color.rgb = BORDER
        ln.line.width = Pt(1.0)
    items = s.get("items", [])[:2]
    y = 1.6
    for i, it in enumerate(items):
        mk = shape(slide, MSO_SHAPE.OVAL, 8.7, y, 0.4, 0.4, fill=None, line_color=O_MID, line_w=2.0)
        label_in(mk, str(i + 1), 10, O_MID)
        text(slide, 9.3, y, 3.2, 0.3, it.get("head", ""), 11, XBOLD, INK)
        text(slide, 9.3, y + 0.35, 3.2, 1.6, it.get("body", ""), 8.5, BODY, GRAY, line=1.35)
        y += 2.3


def lay_cards3(prs, slide, spec, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
    header(slide, spec, s)
    text(slide, 0.9, 1.25, 6.0, 0.6, s.get("title", ""), 22, XBOLD, INK)
    if s.get("body"):
        text(slide, 7.2, 1.3, 5.2, 0.7, s["body"], 8, BODY, GRAY, line=1.3)
    items = s.get("items", [])[:3]
    n = max(len(items), 1)
    cw = (SW - 1.8 - 0.4 * (n - 1)) / n
    for i, it in enumerate(items):
        x = 0.9 + i * (cw + 0.4)
        card(slide, x, 2.3, cw, 4.4)
        photo(slide, it.get("image"), x + 0.25, 2.55, cw - 0.5, 2.2)
        gcirc = grad_rect(slide, x + cw / 2 - 0.35, 4.5, 0.7, 0.7, round_=True, adj=0.5, angle=45.0)
        label_in(gcirc, it.get("icon", "●"), 13)
        text(slide, x + 0.2, 5.35, cw - 0.4, 0.3, it.get("head", ""), 11.5, XBOLD, INK,
             align=PP_ALIGN.CENTER)
        text(slide, x + 0.3, 5.7, cw - 0.6, 0.85, it.get("body", ""), 7.5, BODY, GRAY,
             align=PP_ALIGN.CENTER, line=1.25)


def lay_gradsection(prs, slide, spec, s):
    grad_rect(slide, 0, 0, SW, SH, angle=55.0)
    rings(slide, 11.9, 1.4, 1.5, 2, RGBColor(0xF4, 0x9B, 0x5A))
    header(slide, spec, s, dark=True)
    text(slide, 0.9, 1.5, 6.2, 1.0, s.get("title", ""), 26, XBOLD, WHITE, line=1.1)
    if s.get("kicker"):
        text(slide, 0.9, 2.55, 6.0, 0.3, s["kicker"], 10, SEMI, RGBColor(0xFF, 0xE3, 0xC9))
    if s.get("body"):
        text(slide, 0.9, 3.05, 5.8, 3.3, s["body"], 9.5, BODY, WHITE, line=1.5)
    ph = photo(slide, s.get("image"), 7.3, 1.6, 5.1, 4.6)
    shape(slide, MSO_SHAPE.RECTANGLE, 7.3, 1.6, 5.1, 4.6, fill=None, line_color=WHITE, line_w=1.5)


def lay_thanks(prs, slide, spec, s):
    grad_rect(slide, 0, 0, SW, SH, angle=55.0)
    rings(slide, SW / 2, SH / 2, 2.8, 3, RGBColor(0xF4, 0x9B, 0x5A))
    rings(slide, 1.6, 6.2, 1.6, 2, RGBColor(0xEE, 0x7A, 0x3A))
    text(slide, 1.0, 1.7, SW - 2.0, 0.4, spec.get("brand", ""), 11, XBOLD, WHITE,
         align=PP_ALIGN.CENTER, spacing=300, upper=True)
    text(slide, 1.0, 3.0, SW - 2.0, 1.2, s.get("title", "THANKS"), 54, XBOLD, WHITE,
         align=PP_ALIGN.CENTER, spacing=200, upper=True)
    if s.get("sub"):
        text(slide, 1.0, 4.45, SW - 2.0, 0.4, s["sub"], 11, SEMI, WHITE, align=PP_ALIGN.CENTER, spacing=150)
    meta = s.get("meta", [])[:3]
    if meta:
        cw = (SW - 3.0) / len(meta)
        for i, (k, v) in enumerate(meta):
            x = 1.5 + i * cw
            text(slide, x, 5.6, cw - 0.2, 0.3, k, 9, XBOLD, WHITE, align=PP_ALIGN.CENTER)
            text(slide, x, 5.92, cw - 0.2, 0.3, v, 8.5, BODY, RGBColor(0xFF, 0xE3, 0xC9),
                 align=PP_ALIGN.CENTER)


LAYOUTS = {
    "cover": lay_cover, "contents": lay_contents, "split": lay_split, "project": lay_project,
    "teamrow": lay_teamrow, "stats": lay_stats, "bars": lay_bars, "cards3": lay_cards3,
    "gradsection": lay_gradsection, "thanks": lay_thanks,
}


def build(spec_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation()
    prs.slide_width = Emu(int(SW * 914400))
    prs.slide_height = Emu(int(SH * 914400))
    blank = prs.slide_layouts[6]
    for i, s in enumerate(spec["slides"]):
        s["_page"] = i + 1
        slide = prs.slides.add_slide(blank)
        fn = LAYOUTS.get(s.get("layout"))
        if not fn:
            raise SystemExit(f"unknown layout: {s.get('layout')}")
        fn(prs, slide, spec, s)
    out = spec.get("output", "deck.pptx")
    prs.save(out)
    print(f"OK {out} ({len(spec['slides'])} slides)")


if __name__ == "__main__":
    build(sys.argv[1])
