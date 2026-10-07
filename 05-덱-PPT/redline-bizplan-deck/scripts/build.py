# redline-bizplan-deck builder: spec JSON -> red/navy business plan pptx
# usage: python build.py spec.json
import json, os, sys, tempfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageOps

RED = RGBColor(0xD9, 0x33, 0x33)
NAVY = RGBColor(0x2F, 0x32, 0x44)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x7A, 0x7A, 0x80)
LIGHT = RGBColor(0xF5, 0xF5, 0xF6)
PLACEHOLDER = RGBColor(0xDE, 0xDE, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
XBOLD = "Pretendard ExtraBold"
SEMI = "Pretendard SemiBold"
BODY = "Pretendard Regular"
SW, SH = 13.333, 7.5
_TMP = tempfile.mkdtemp(prefix="rbd_")


def _fonts(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def text(slide, x, y, w, h, s, size=10, font=BODY, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.2, wrap=True, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)
        if spacing is not None:
            r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def twotone(slide, x, y, w, h, s, size=36, base=INK, accent=RED, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, line=1.02, font=XBOLD):
    """Title with '|' marking the accent part; supports \\n for multiple lines."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        parts = ln.split("|")
        for j, part in enumerate(parts):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            r.font.size = Pt(size)
            r.font.color.rgb = base if j == 0 else accent
            _fonts(r, font)
    return tb


def shape(slide, kind, x, y, w, h, fill=RED, line_color=None, line_w=1.0):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    return sp


def label_in(sp, s, size=12, color=WHITE, font=XBOLD):
    tf = sp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)


def circle(slide, cx, cy, d, fill=RED, line_color=None, line_w=1.5, lbl=None, size=None, color=WHITE):
    sp = shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill, line_color, line_w)
    if lbl is not None:
        label_in(sp, lbl, size or max(9, d * 18), color)
    return sp


def hline(slide, x, y, w, weight=1.0, color=RGBColor(0xE2, 0xE2, 0xE6)):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def photo(slide, path, x, y, w, h, dark=False):
    """Color center-crop fill; light-gray placeholder if missing."""
    if path and os.path.isfile(path):
        im = Image.open(path)
        im = ImageOps.exif_transpose(im).convert("RGB")
        tw, th = max(int(w * 150), 50), max(int(h * 150), 50)
        im = ImageOps.fit(im, (tw, th), Image.LANCZOS)
        out = os.path.join(_TMP, f"p{abs(hash((path, tw, th)))}.jpg")
        im.save(out, quality=90)
        slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h,
              fill=(RGBColor(0x3C, 0x40, 0x55) if dark else PLACEHOLDER))


def slash(slide, x, y, w, h, fill=RED):
    """Red diagonal parallelogram accent."""
    sp = shape(slide, MSO_SHAPE.PARALLELOGRAM, x, y, w, h, fill=fill)
    try:
        sp.adjustments[0] = 0.55
    except Exception:
        pass
    return sp


def header(slide, s, dark=False):
    base = WHITE if dark else INK
    twotone(slide, 0.55, 0.35, 9.5, 0.85, s.get("title", ""), 30, base=base)
    if s.get("sub"):
        text(slide, 0.55, 1.12, 9.0, 0.3, s["sub"], 9.5, BODY, GRAY)
    slash(slide, 11.9, 0.45, 0.55, 0.28)
    slash(slide, 12.55, 0.45, 0.3, 0.28, fill=NAVY)
    n = s.get("_page")
    if n:
        text(slide, SW - 0.9, SH - 0.42, 0.5, 0.3, f"{n:02d}", 9, SEMI, GRAY, align=PP_ALIGN.RIGHT)


# ---------------- layouts ----------------

def lay_cover(prs, slide, s):
    if s.get("logo_text"):
        circle(slide, 0.8, 0.75, 0.42, fill=None, line_color=RED, line_w=2.0)
        text(slide, 1.15, 0.6, 3.5, 0.35, s["logo_text"], 11, SEMI, INK)
    twotone(slide, 0.55, 2.0, 6.6, 2.4, s.get("title", "BUSINESS\nPLAN"), 54)
    if s.get("tagline"):
        text(slide, 0.55, 4.35, 5.8, 0.35, "✦ " + s["tagline"], 11, SEMI, RED)
    if s.get("body"):
        text(slide, 0.55, 4.85, 5.6, 1.0, s["body"], 9.5, BODY, GRAY, line=1.3)
    # barcode-ish deco
    for i, bw in enumerate([0.04, 0.02, 0.05, 0.02, 0.03, 0.05, 0.02, 0.04, 0.03, 0.05,
                            0.02, 0.03, 0.04, 0.02, 0.05, 0.03]):
        shape(slide, MSO_SHAPE.RECTANGLE, 0.55 + i * 0.11, 6.35, bw, 0.5, fill=INK)
    if s.get("meta"):
        text(slide, 0.55, 7.0, 6.0, 0.3, s["meta"], 8.5, BODY, GRAY)
    # photo right with red year block
    photo(slide, s.get("image"), 6.9, 1.15, 5.9, 5.2)
    blk = slash(slide, 6.35, 1.15, 1.7, 5.2)
    if s.get("year"):
        tf = blk.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(s["year"])
        r.font.size = Pt(20)
        r.font.color.rgb = WHITE
        _fonts(r, XBOLD)
        blk.rotation = 0


def lay_welcome(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 7.3, 1.6, 5.45, 5.35)
    slash(slide, 6.9, 1.6, 1.0, 1.6)
    if s.get("name"):
        text(slide, 0.55, 1.9, 5.8, 0.35, s["name"], 14, XBOLD, INK)
        text(slide, 0.55, 2.28, 5.8, 0.3, s.get("role", ""), 9.5, BODY, RED)
    if s.get("message_head"):
        circle(slide, 0.85, 3.25, 0.55, fill=RED, lbl=s.get("icon", "💬"), size=13)
        text(slide, 1.35, 3.1, 5.0, 0.3, s["message_head"], 11, XBOLD, INK)
    text(slide, 0.55, 3.75, 6.1, 3.2, s.get("body", ""), 9.5, BODY, GRAY, line=1.4)


def lay_toc(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 7.0, 1.9, 5.75, 4.4)
    slash(slide, 6.6, 5.5, 1.2, 0.8)
    if s.get("body"):
        text(slide, 0.55, 1.75, 5.7, 0.9, s["body"], 9, BODY, GRAY, line=1.3)
    items = s.get("items", [])[:8]
    y = 2.75
    for i, it in enumerate(items):
        text(slide, 0.55, y, 0.7, 0.3, f"{i+1:02d}.", 11, XBOLD, RED)
        text(slide, 1.25, y, 5.2, 0.3, it, 11, SEMI, INK)
        y += 0.52


def lay_about(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 8.6, 1.6, 4.2, 5.3)
    slash(slide, 8.2, 1.6, 1.0, 1.4, fill=NAVY)
    if s.get("body"):
        text(slide, 0.55, 1.75, 7.4, 1.1, s["body"], 9.5, BODY, GRAY, line=1.35)
    items = s.get("items", [])[:3]
    xw = 7.6 / max(len(items), 1)
    for i, it in enumerate(items):
        x = 0.55 + i * xw
        c = NAVY if i % 2 == 0 else RED
        circle(slide, x + 0.45, 3.55, 0.9, fill=c, lbl=it.get("icon", f"{i+1:02d}"), size=16)
        text(slide, x, 4.25, xw - 0.4, 0.3, f"{i+1:02d}", 10, XBOLD, RED)
        text(slide, x, 4.55, xw - 0.4, 0.35, it.get("head", ""), 11.5, XBOLD, INK)
        text(slide, x, 4.95, xw - 0.4, 1.8, it.get("body", ""), 8.5, BODY, GRAY, line=1.3)


def lay_split(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 7.9, 1.6, 4.9, 4.1)
    blk = slash(slide, 7.35, 1.6, 1.5, 4.1)
    if s.get("badge"):
        label_in(blk, s["badge"], 18)
    cols = s.get("cols", [])[:2]
    cw = 3.3
    for i, c in enumerate(cols):
        x = 0.55 + i * (cw + 0.25)
        hline(slide, x, 1.85, cw, 2.0, RED if i == 0 else NAVY)
        text(slide, x, 2.05, cw, 4.3, c, 9.5, BODY, GRAY, line=1.4)
    if s.get("footer"):
        text(slide, 0.55, 6.55, 12.2, 0.5, s["footer"], 9, SEMI, INK)


def lay_darkcards(prs, slide, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=NAVY)
    slash(slide, 0.55, 0.85, 0.9, 0.45)
    twotone(slide, 0.55, 0.75, 12.2, 0.8, s.get("title", ""), 32, base=WHITE, align=PP_ALIGN.CENTER)
    if s.get("sub"):
        text(slide, 0.55, 1.6, 12.2, 0.3, s["sub"], 9.5, BODY, RGBColor(0xC9, 0xCB, 0xD6),
             align=PP_ALIGN.CENTER)
    items = s.get("items", [])[:3]
    n = max(len(items), 1)
    gap = 0.4
    cw = (SW - 1.6 - gap * (n - 1)) / n
    for i, it in enumerate(items):
        x = 0.8 + i * (cw + gap)
        shape(slide, MSO_SHAPE.RECTANGLE, x, 2.75, cw, 3.9, fill=RGBColor(0x3A, 0x3E, 0x53))
        circle(slide, x + cw / 2, 3.6, 0.85, fill=RED, lbl=it.get("icon", "✦"), size=17)
        text(slide, x + 0.2, 4.25, cw - 0.4, 0.35, it.get("head", ""), 12, XBOLD, WHITE,
             align=PP_ALIGN.CENTER)
        text(slide, x + 0.25, 4.7, cw - 0.5, 1.7, it.get("body", ""), 8.5, BODY,
             RGBColor(0xC9, 0xCB, 0xD6), align=PP_ALIGN.CENTER, line=1.3)
    n2 = s.get("_page")
    if n2:
        text(slide, SW - 0.9, SH - 0.42, 0.5, 0.3, f"{n2:02d}", 9, SEMI,
             RGBColor(0x9A, 0x9D, 0xB0), align=PP_ALIGN.RIGHT)


def lay_problems(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 0.55, 4.6, 5.6, 2.3)
    if s.get("body"):
        text(slide, 0.55, 1.75, 5.6, 2.6, s["body"], 9.5, BODY, GRAY, line=1.4)
    panel_x = 6.7
    shape(slide, MSO_SHAPE.RECTANGLE, panel_x, 1.5, SW - panel_x - 0.0, 6.0, fill=NAVY)
    slash(slide, panel_x - 0.5, 1.5, 1.1, 1.1)
    if s.get("panel_head"):
        text(slide, panel_x + 0.5, 1.85, 5.6, 0.35, s["panel_head"], 13, XBOLD, WHITE)
    items = s.get("items", [])[:3]
    y = 2.6
    for i, it in enumerate(items):
        circle(slide, panel_x + 0.85, y + 0.3, 0.6, fill=None, line_color=RED, line_w=2.0,
               lbl=f"{i+1:02d}", size=13, color=WHITE)
        text(slide, panel_x + 1.5, y + 0.02, 4.7, 0.35, it.get("head", ""), 11.5, XBOLD, WHITE)
        text(slide, panel_x + 1.5, y + 0.4, 4.7, 1.0, it.get("body", ""), 8.5, BODY,
             RGBColor(0xC9, 0xCB, 0xD6), line=1.3)
        y += 1.5


def lay_numphoto(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 0.55, 1.9, 5.6, 4.6)
    slash(slide, 5.7, 1.9, 1.0, 1.2)
    if s.get("body"):
        text(slide, 7.0, 1.8, 5.6, 1.0, s["body"], 9, BODY, GRAY, line=1.3)
    items = s.get("items", [])[:3]
    y = 3.0
    for i, it in enumerate(items):
        circle(slide, 7.35, y + 0.3, 0.6, fill=None, line_color=RED, line_w=2.0,
               lbl=f"{i+1:02d}", size=13, color=INK)
        text(slide, 8.0, y + 0.02, 4.6, 0.35, it.get("head", ""), 11.5, XBOLD, INK)
        text(slide, 8.0, y + 0.4, 4.6, 0.85, it.get("body", ""), 8.5, BODY, GRAY, line=1.3)
        hline(slide, 7.0, y + 1.18, 5.6, 1.0)
        y += 1.35


def lay_progress(prs, slide, s):
    header(slide, s)
    if s.get("body"):
        text(slide, 0.55, 1.7, 12.0, 0.6, s["body"], 9.5, BODY, GRAY, line=1.3)
    items = s.get("items", [])[:3]
    n = max(len(items), 1)
    cw = (SW - 1.1) / n
    cy = 3.55
    hline(slide, 0.55 + cw / 2, cy, cw * (n - 1), 1.5, RGBColor(0xE2, 0xE2, 0xE6))
    for i, it in enumerate(items):
        cx = 0.55 + cw * i + cw / 2
        c = RED if i % 2 == 0 else NAVY
        circle(slide, cx, cy, 0.95, fill=c, lbl=it.get("head", f"{i+1:02d}"), size=10)
        text(slide, cx - cw / 2 + 0.3, 4.25, cw - 0.6, 2.4, it.get("body", ""), 8.5, BODY, GRAY,
             align=PP_ALIGN.CENTER, line=1.3)


def lay_team(prs, slide, s):
    header(slide, s)
    if s.get("body"):
        text(slide, 0.55, 1.75, 3.4, 4.5, s["body"], 9, BODY, GRAY, line=1.4)
    items = s.get("items", [])[:6]
    cols = 3
    x0 = 4.3 if s.get("body") else 0.8
    gw = (SW - x0 - 0.55)
    cw = gw / cols
    for i, it in enumerate(items):
        r_, c_ = divmod(i, cols)
        x = x0 + c_ * cw
        y = 1.6 + r_ * 2.85
        photo(slide, it.get("image"), x + 0.2, y, cw - 0.55, 1.85)
        text(slide, x + 0.2, y + 1.95, cw - 0.55, 0.3, it.get("name", ""), 10.5, XBOLD, INK)
        text(slide, x + 0.2, y + 2.25, cw - 0.55, 0.3, it.get("role", ""), 8.5, BODY, RED)


def lay_skillbars(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 0.55, 1.75, 4.3, 4.4)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.55, 5.35, 4.3, 0.8, fill=NAVY)
    text(slide, 0.75, 5.5, 3.9, 0.3, s.get("name", ""), 12, XBOLD, WHITE)
    text(slide, 0.75, 5.82, 3.9, 0.25, s.get("role", ""), 8.5, BODY, RGBColor(0xC9, 0xCB, 0xD6))
    if s.get("body"):
        text(slide, 5.4, 1.8, 7.2, 1.4, s["body"], 9, BODY, GRAY, line=1.35)
    bars = s.get("bars", [])[:4]
    y = 3.5
    for i, b in enumerate(bars):
        pct = max(0, min(int(b.get("pct", 50)), 100))
        text(slide, 5.4, y, 4.5, 0.28, b.get("head", ""), 10, SEMI, INK)
        text(slide, 11.9, y, 0.9, 0.28, f"{pct}%", 10, XBOLD, RED, align=PP_ALIGN.RIGHT)
        shape(slide, MSO_SHAPE.RECTANGLE, 5.4, y + 0.35, 7.4, 0.12, fill=LIGHT)
        shape(slide, MSO_SHAPE.RECTANGLE, 5.4, y + 0.35, 7.4 * pct / 100.0, 0.12,
              fill=RED if i % 2 == 0 else NAVY)
        y += 0.8


def lay_services(prs, slide, s):
    header(slide, s)
    photo(slide, s.get("image"), 0.55, 4.4, 5.2, 2.7)
    if s.get("body"):
        text(slide, 0.55, 1.8, 5.2, 2.4, s["body"], 9.5, BODY, GRAY, line=1.4)
    items = s.get("items", [])[:4]
    for i, it in enumerate(items):
        r_, c_ = divmod(i, 2)
        x = 6.5 + c_ * 3.3
        y = 1.9 + r_ * 2.6
        c = RED if (r_ + c_) % 2 == 0 else NAVY
        circle(slide, x + 0.5, y + 0.5, 1.0, fill=c, lbl=it.get("icon", "✦"), size=20)
        text(slide, x, y + 1.15, 2.9, 0.3, it.get("head", ""), 11, XBOLD, INK)
        text(slide, x, y + 1.48, 2.9, 0.9, it.get("body", ""), 8.5, BODY, GRAY, line=1.25)


def lay_contact(prs, slide, s):
    twotone(slide, 0.55, 1.0, 12.2, 1.0, s.get("title", "CONTACT |ME"), 44)
    slash(slide, 11.9, 1.1, 0.55, 0.28)
    slash(slide, 12.55, 1.1, 0.3, 0.28, fill=NAVY)
    if s.get("body"):
        text(slide, 0.55, 2.15, 8.0, 0.8, s["body"], 9.5, BODY, GRAY, line=1.35)
    items = s.get("items", [])[:3]
    n = max(len(items), 1)
    cw = (SW - 1.1) / n
    for i, it in enumerate(items):
        x = 0.55 + i * cw
        c = RED if i % 2 == 0 else NAVY
        circle(slide, x + cw / 2, 4.3, 1.0, fill=c, lbl=it.get("icon", "✦"), size=20)
        text(slide, x, 5.05, cw, 0.3, it.get("head", ""), 11, XBOLD, INK, align=PP_ALIGN.CENTER)
        text(slide, x, 5.4, cw, 0.6, it.get("body", ""), 9, BODY, GRAY, align=PP_ALIGN.CENTER, line=1.25)
    n2 = s.get("_page")
    if n2:
        text(slide, SW - 0.9, SH - 0.42, 0.5, 0.3, f"{n2:02d}", 9, SEMI, GRAY, align=PP_ALIGN.RIGHT)


LAYOUTS = {
    "cover": lay_cover, "welcome": lay_welcome, "toc": lay_toc, "about": lay_about,
    "split": lay_split, "darkcards": lay_darkcards, "problems": lay_problems,
    "numphoto": lay_numphoto, "progress": lay_progress, "team": lay_team,
    "skillbars": lay_skillbars, "services": lay_services, "contact": lay_contact,
}


def build(spec_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation()
    prs.slide_width = Emu(int(SW * 914400))
    prs.slide_height = Emu(int(SH * 914400))
    blank = prs.slide_layouts[6]
    for i, s in enumerate(spec["slides"]):
        s["_page"] = i + 1
        slide = prs.slides.add_slide(blank)
        shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
        fn = LAYOUTS.get(s.get("layout"))
        if not fn:
            raise SystemExit(f"unknown layout: {s.get('layout')}")
        fn(prs, slide, s)
    out = spec.get("output", "deck.pptx")
    prs.save(out)
    print(f"OK {out} ({len(spec['slides'])} slides)")


if __name__ == "__main__":
    build(sys.argv[1])
