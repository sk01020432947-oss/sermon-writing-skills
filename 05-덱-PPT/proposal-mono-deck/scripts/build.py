# proposal-mono-deck builder: spec JSON -> framed B/W project proposal pptx
# usage: python build.py spec.json
import json, os, sys, tempfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageOps

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x11, 0x11, 0x11)
DARK = RGBColor(0x14, 0x14, 0x14)
GRAY = RGBColor(0x8A, 0x8A, 0x8A)
GHOST = RGBColor(0xDC, 0xDC, 0xDC)
TINT = RGBColor(0xE9, 0xE9, 0xE9)
BAND = RGBColor(0x9A, 0x9A, 0x9A)
PLACEHOLDER = RGBColor(0xD6, 0xD6, 0xD6)
XBOLD = "Pretendard ExtraBold"
BLACK = "Pretendard Black"
SEMI = "Pretendard SemiBold"
BODY = "Pretendard Regular"
SW, SH = 13.333, 7.5
M = 0.32  # frame margin
_TMP = tempfile.mkdtemp(prefix="pmd_")


def _fonts(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def text(slide, x, y, w, h, s, size=9, font=BODY, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.25, wrap=True, spacing=None, upper=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run()
        r.text = ln.upper() if upper else ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)
        if spacing is not None:
            r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def rect(slide, x, y, w, h, fill=INK, line_color=None, line_w=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    return sp


def hline(slide, x, y, w, weight=0.75, color=INK):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def photo(slide, path, x, y, w, h, dark=False):
    if path and os.path.isfile(path):
        im = Image.open(path)
        im = ImageOps.exif_transpose(im).convert("L")
        im = ImageOps.autocontrast(im, cutoff=1)
        tw, th = max(int(w * 150), 50), max(int(h * 150), 50)
        im = ImageOps.fit(im, (tw, th), Image.LANCZOS)
        out = os.path.join(_TMP, f"p{abs(hash((path, tw, th)))}.jpg")
        im.convert("RGB").save(out, quality=90)
        slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        rect(slide, x, y, w, h, fill=(RGBColor(0x3A, 0x3A, 0x3A) if dark else PLACEHOLDER))


def ghost(slide, x, y, s, size=54, align=PP_ALIGN.LEFT, color=GHOST):
    text(slide, x, y, 2.5, 1.1, s, size, BLACK, color, align=align)


def chrome(slide, spec, s, dark=False):
    """Outer frame + top meta row, on every slide."""
    fg = WHITE if dark else INK
    rect(slide, M, M, SW - 2 * M, SH - 2 * M, fill=None, line_color=fg, line_w=1.0)
    y = M + 0.14
    text(slide, M + 0.25, y, 2.5, 0.35, "Project\n" + spec.get("project", "Name"), 6.5, SEMI, fg, line=1.1)
    text(slide, SW / 2 - 1.25, y, 2.5, 0.35, "Date:\n" + spec.get("date", ""), 6.5, SEMI, fg,
         align=PP_ALIGN.CENTER, line=1.1)
    text(slide, SW - M - 2.75, y, 2.5, 0.35, "Company\n" + spec.get("company", "Name"), 6.5, SEMI, fg,
         align=PP_ALIGN.RIGHT, line=1.1)
    hline(slide, M + 0.25, y + 0.5, SW - 2 * M - 0.5, 0.75, fg)


# ---------------- layouts ----------------

def lay_cover(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=DARK)
    chrome(slide, spec, s, dark=True)
    photo(slide, s.get("image"), 8.35, 1.35, 3.9, 4.9, dark=True)
    text(slide, M + 0.35, 1.5, 7.2, 2.0, s.get("title", "PROJECT\nPROPOSAL."), 44, BLACK, WHITE,
         line=1.0, upper=True)
    band = rect(slide, M, 3.7, 7.6, 1.35, fill=BAND)
    if s.get("body"):
        text(slide, M + 0.35, 3.9, 6.9, 1.0, s["body"], 8, BODY, RGBColor(0x2A, 0x2A, 0x2A), line=1.3)
    if s.get("credit"):
        hline(slide, M + 0.35, 6.2, 1.0, 0.75, WHITE)
        text(slide, M + 1.5, 6.05, 5.5, 0.5, s["credit"], 8.5, SEMI, WHITE, line=1.2)


def lay_index(prs, slide, spec, s):
    chrome(slide, spec, s)
    text(slide, M + 0.35, 1.4, 5.5, 1.5, s.get("title", "INDEX\nCONTENT"), 34, BLACK, INK, line=1.0, upper=True)
    if s.get("sub"):
        text(slide, M + 0.35, 3.15, 4.8, 0.3, s["sub"], 10, SEMI, INK)
    if s.get("body"):
        text(slide, M + 0.35, 3.55, 4.6, 2.5, s["body"], 8.5, BODY, GRAY, line=1.35)
    items = s.get("items", [])[:8]
    x0, w = 7.0, SW - 7.0 - M - 0.35
    rect(slide, 6.7, 1.2, SW - 6.7 - M - 0.15, 0.35 + len(items) * 0.55, fill=TINT)
    y = 1.45
    for i, it in enumerate(items):
        text(slide, x0, y, w - 0.7, 0.3, f"{i+1}. {it}", 9.5, SEMI, INK)
        text(slide, x0 + w - 0.6, y, 0.5, 0.3, f"{i+2:02d}", 8.5, BODY, GRAY, align=PP_ALIGN.RIGHT)
        hline(slide, x0, y + 0.32, w - 0.1, 0.75, INK)
        y += 0.55


def lay_section(prs, slide, spec, s):
    chrome(slide, spec, s)
    text(slide, M + 0.35, 1.35, 5.8, 1.6, s.get("title", ""), 34, BLACK, INK, line=1.0, upper=True)
    if s.get("number"):
        ghost(slide, M + 0.4, 2.85, s["number"], 44)
    if s.get("body"):
        text(slide, 2.1, 3.15, 3.9, 2.2, s["body"], 8.5, BODY, GRAY, line=1.4)
    if s.get("sub"):
        text(slide, M + 0.35, 5.6, 3.5, 0.3, s["sub"], 9.5, SEMI, INK)
        text(slide, M + 0.35, 5.95, 3.9, 0.9, s.get("footnote", ""), 7.5, BODY, GRAY, line=1.3)
    photo(slide, s.get("image"), 6.6, 1.35, 5.0, 4.4)
    rect(slide, 6.2, 4.6, 2.1, 2.1, fill=TINT)
    photo_note = s.get("caption")
    if photo_note:
        text(slide, 8.5, 6.0, 4.0, 0.6, photo_note, 7.5, BODY, GRAY, line=1.25)


def lay_detail(prs, slide, spec, s):
    chrome(slide, spec, s)
    text(slide, M + 0.35, 1.35, 5.4, 1.5, s.get("title", ""), 32, BLACK, INK, line=1.0, upper=True)
    if s.get("sub"):
        text(slide, M + 0.35, 3.0, 2.2, 0.8, s["sub"], 9.5, SEMI, INK, line=1.25)
    if s.get("body"):
        text(slide, 8.3, 1.45, 4.2, 2.4, s["body"], 8, BODY, GRAY, line=1.35)
    photo(slide, s.get("image"), 4.35, 2.5, 4.7, 3.6)
    if s.get("body2"):
        text(slide, 9.3, 4.3, 3.2, 2.0, s["body2"], 8, BODY, GRAY, line=1.35)
    if s.get("number"):
        ghost(slide, M + 0.4, 5.7, s["number"], 48)


def lay_framed(prs, slide, spec, s):
    chrome(slide, spec, s)
    text(slide, M + 0.35, 1.35, 5.6, 1.5, s.get("title", ""), 32, BLACK, INK, line=1.0, upper=True)
    if s.get("sub"):
        text(slide, M + 0.35, 3.05, 2.4, 0.8, s["sub"], 9.5, SEMI, INK, line=1.25)
    if s.get("body"):
        text(slide, M + 0.35, 4.0, 4.6, 2.4, s["body"], 8.5, BODY, GRAY, line=1.4)
    fx, fy, fw, fh = 6.8, 1.5, 4.3, 3.4
    photo(slide, s.get("image"), fx, fy, fw, fh)
    rect(slide, fx - 0.25, fy - 0.25, fw + 0.5, fh + 0.5, fill=None, line_color=INK, line_w=1.25)
    if s.get("caption"):
        vt = rect(slide, fx + fw + 0.42, fy - 0.25, 0.45, fh + 0.5, fill=None)
        tf = vt.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = s["caption"]
        r.font.size = Pt(7.5)
        r.font.color.rgb = GRAY
        _fonts(r, SEMI)
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        bodyPr.set("vert", "vert270")
    if s.get("number"):
        ghost(slide, M + 0.4, 5.75, s["number"], 44)


def lay_stages(prs, slide, spec, s):
    chrome(slide, spec, s)
    text(slide, M + 0.35, 1.35, 5.4, 1.5, s.get("title", ""), 32, BLACK, INK, line=1.0, upper=True)
    if s.get("body"):
        text(slide, 6.3, 1.45, 3.2, 2.6, s["body"], 8, BODY, GRAY, line=1.35)
    if s.get("body2"):
        text(slide, 9.8, 1.45, 2.7, 2.6, s["body2"], 8, BODY, GRAY, line=1.35)
    items = s.get("items", [])[:3]
    n = max(len(items), 1)
    x0 = M + 0.35
    tw = SW - 2 * (M + 0.35)
    cw = tw / n
    for i, it in enumerate(items):
        x = x0 + i * cw
        photo(slide, it.get("image"), x, 4.15, 1.55, 1.25)
        text(slide, x + 1.75, 4.15, cw - 2.0, 0.35, it.get("head", ""), 9.5, SEMI, INK, line=1.15)
        text(slide, x + 1.75, 4.55, cw - 2.0, 1.6, it.get("body", ""), 7.5, BODY, GRAY, line=1.3)


def lay_team(prs, slide, spec, s):
    chrome(slide, spec, s)
    text(slide, M + 0.35, 1.35, 5.0, 1.5, s.get("title", "OUR\nTEAM"), 34, BLACK, INK, line=1.0, upper=True)
    if s.get("sub"):
        text(slide, M + 0.35, 3.2, 3.0, 0.3, s["sub"], 9.5, SEMI, INK)
    if s.get("body"):
        text(slide, M + 0.35, 3.6, 4.2, 1.6, s["body"], 8, BODY, GRAY, line=1.35)
    if s.get("number"):
        ghost(slide, M + 0.4, 5.5, s["number"], 48)
    items = s.get("items", [])[:3]
    y = 1.35
    for i, it in enumerate(items):
        photo(slide, it.get("image"), 5.9, y, 1.6, 1.6)
        text(slide, 7.75, y + 0.05, 2.2, 0.3, it.get("name", ""), 9.5, SEMI, INK)
        text(slide, 7.75, y + 0.38, 4.8, 1.2, it.get("body", ""), 7.5, BODY, GRAY, line=1.3)
        y += 1.85


def lay_feature(prs, slide, spec, s):
    chrome(slide, spec, s)
    left = s.get("photo_side", "left") == "left"
    px = M + 0.15 if left else 6.9
    tx = 6.5 if left else M + 0.35
    photo(slide, s.get("image"), px, 1.3, 5.9, 5.4)
    rect(slide, px + (4.3 if left else -0.5), 0.95, 2.1, 2.1, fill=TINT)
    text(slide, tx, 2.2, 6.0, 1.5, s.get("title", ""), 34, BLACK, INK, line=1.0, upper=True)
    if s.get("sub"):
        text(slide, tx, 3.9, 2.4, 0.8, s["sub"], 9.5, SEMI, INK, line=1.25)
    if s.get("body"):
        text(slide, tx, 4.75, 5.6, 1.6, s["body"], 8.5, BODY, GRAY, line=1.4)
    if s.get("note"):
        nb = rect(slide, tx, 6.0, 4.6, 0.75, fill=None, line_color=INK, line_w=0.75)
        text(slide, tx + 0.2, 6.15, 4.2, 0.5, s["note"], 7.5, BODY, GRAY, line=1.25)
    if s.get("number"):
        ghost(slide, tx + 4.9, 3.75, s["number"], 44)


def lay_timeline(prs, slide, spec, s):
    chrome(slide, spec, s)
    text(slide, M + 0.35, 1.4, 11.0, 1.0, s.get("title", "PROJECT TIMELINE"), 34, BLACK, INK, upper=True)
    items = s.get("items", [])[:6]
    n = max(len(items), 1)
    x0 = M + 0.7
    tw = SW - 2 * (M + 0.7)
    cy = 4.35
    hline(slide, x0 - 0.2, cy, tw + 0.4, 1.0, INK)
    for i, it in enumerate(items):
        cx = x0 + (tw * i / max(n - 1, 1)) if n > 1 else x0 + tw / 2
        d = 0.12
        rect(slide, cx - d / 2, cy - d / 2, d, d, fill=INK)
        up = i % 2 == 0
        ty = cy - 1.75 if up else cy + 0.25
        tx = min(max(cx - 1.0, M + 0.25), SW - M - 0.25 - 2.0)  # clamp inside frame
        text(slide, tx, ty, 2.0, 0.35, it.get("head", ""), 13, XBOLD, INK, align=PP_ALIGN.CENTER)
        text(slide, tx, ty + 0.38, 2.0, 0.3, it.get("sub", ""), 8, SEMI, INK, align=PP_ALIGN.CENTER)
        text(slide, tx, ty + 0.68, 2.0, 0.85, it.get("body", ""), 7, BODY, GRAY,
             align=PP_ALIGN.CENTER, line=1.25)


def lay_closing(prs, slide, spec, s):
    chrome(slide, spec, s)
    photo(slide, s.get("image"), 7.5, 1.3, 5.0, 5.4)
    text(slide, M + 0.35, 1.6, 6.8, 2.0, s.get("title", "THANK YOU FOR\nVIEWING."), 34, BLACK, INK,
         line=1.05, upper=True)
    y = 4.0
    for k, v in s.get("meta", [])[:5]:
        text(slide, M + 0.35, y, 1.6, 0.3, k, 8.5, SEMI, INK)
        text(slide, 2.3, y, 4.6, 0.3, v, 8.5, BODY, GRAY)
        y += 0.42
    if s.get("body"):
        text(slide, M + 0.35, y + 0.2, 6.2, 1.2, s["body"], 8, BODY, GRAY, line=1.35)


LAYOUTS = {
    "cover": lay_cover, "index": lay_index, "section": lay_section, "detail": lay_detail,
    "framed": lay_framed, "stages": lay_stages, "team": lay_team, "feature": lay_feature,
    "timeline": lay_timeline, "closing": lay_closing,
}


def build(spec_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation()
    prs.slide_width = Emu(int(SW * 914400))
    prs.slide_height = Emu(int(SH * 914400))
    blank = prs.slide_layouts[6]
    for i, s in enumerate(spec["slides"]):
        s["_page"] = i + 1
        slide = prs.slides.add_slide(blank)
        rect(slide, 0, 0, SW, SH, fill=WHITE)
        fn = LAYOUTS.get(s.get("layout"))
        if not fn:
            raise SystemExit(f"unknown layout: {s.get('layout')}")
        fn(prs, slide, spec, s)
    out = spec.get("output", "deck.pptx")
    prs.save(out)
    print(f"OK {out} ({len(spec['slides'])} slides)")


if __name__ == "__main__":
    build(sys.argv[1])
