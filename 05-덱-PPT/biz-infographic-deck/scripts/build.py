# biz-infographic-deck builder: spec JSON -> business infographic pptx
# usage: python build.py spec.json
import json, math, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

NAVY = RGBColor(0x41, 0x42, 0x63)
PLUM = RGBColor(0x9A, 0x5A, 0x6C)
RED = RGBColor(0xD0, 0x4A, 0x54)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
INK = RGBColor(0x2E, 0x2E, 0x3A)
GRAY = RGBColor(0x8B, 0x8B, 0x94)
LIGHT = RGBColor(0xF2, 0xF2, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CYCLE = [NAVY, PLUM, RED, AMBER]
XBOLD = "Pretendard ExtraBold"
SEMI = "Pretendard SemiBold"
BODY = "Pretendard Regular"
SW, SH = 13.333, 7.5


def _fonts(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def text(slide, x, y, w, h, s, size=10, font=BODY, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.15, wrap=True, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)
        if spacing is not None:
            r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def shape(slide, kind, x, y, w, h, fill=NAVY, line_color=None, line_w=1.0):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    return sp


def label_in(sp, s, size=12, color=WHITE, font=XBOLD):
    tf = sp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)


def circle(slide, cx, cy, d, fill=NAVY, line_color=None, line_w=1.5,
           lbl=None, size=None, color=WHITE):
    sp = shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill, line_color, line_w)
    if lbl is not None:
        label_in(sp, lbl, size or max(9, d * 18), color if fill else (line_color or INK))
    return sp


def hline(slide, x, y, w, weight=1.0, color=LIGHT, dash=None):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    if dash:
        d = ln.line._get_or_add_ln()
        pd = d.makeelement(qn("a:prstDash"), {"val": dash})
        d.append(pd)
    return ln


def pie_wedge(slide, cx, cy, r, pct, color):
    """Progress wedge from 12 o'clock, clockwise, as a freeform polygon (PIE adjustments are unreliable)."""
    pct = max(0.0, min(pct, 100.0))
    if pct <= 0:
        return
    E = 914400
    steps = max(2, int(pct * 0.72))  # ~5° resolution
    pts = []
    for k in range(steps + 1):
        ang = math.radians(-90 + 360.0 * pct / 100.0 * k / steps)
        pts.append((Emu(int((cx + r * math.cos(ang)) * E)), Emu(int((cy + r * math.sin(ang)) * E))))
    fb = slide.shapes.build_freeform(Emu(int(cx * E)), Emu(int(cy * E)), scale=1.0)
    fb.add_line_segments(pts, close=True)
    sp = fb.convert_to_shape()
    sp.shadow.inherit = False
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    return sp


def header(slide, s):
    text(slide, 0.5, 0.28, 9.5, 0.4, s.get("title", ""), 18, XBOLD, INK)
    if s.get("sub"):
        text(slide, 0.5, 0.68, 9.5, 0.3, s["sub"], 9, BODY, GRAY)
    hline(slide, 0.5, 1.05, SW - 1.0, 1.0, LIGHT)
    n = s.get("_page")
    if n:
        text(slide, SW - 1.0, 0.35, 0.6, 0.3, f"{n:02d}", 9, SEMI, GRAY, align=PP_ALIGN.RIGHT)


def col(i):
    return CYCLE[i % len(CYCLE)]


# ---------------- layouts ----------------

def lay_divider(prs, slide, s):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=NAVY)
    # dot band decoration
    for r in range(6):
        for c in range(34):
            if (r * 7 + c * 3) % 5 < 2:
                d = 0.045
                shape(slide, MSO_SHAPE.OVAL, 1.1 + c * 0.33, 1.15 + r * 0.28, d, d,
                      fill=RGBColor(0x5A, 0x5B, 0x7E))
    circle(slide, SW / 2, 2.55, 1.15, fill=None, line_color=WHITE, line_w=1.5,
           lbl=s.get("icon", "◆"), size=26)
    text(slide, 1.0, 3.55, SW - 2.0, 1.1, s.get("title", ""), 40, XBOLD, WHITE,
         align=PP_ALIGN.CENTER, spacing=100)
    if s.get("sub"):
        text(slide, 1.0, 4.55, SW - 2.0, 0.4, s["sub"], 11, BODY, RGBColor(0xB9, 0xBA, 0xD0),
             align=PP_ALIGN.CENTER, spacing=150)
    hline(slide, SW / 2 - 1.5, 5.15, 3.0, 1.0, RGBColor(0x6A, 0x6B, 0x8E))


def lay_steps(prs, slide, s):
    header(slide, s)
    items = s.get("items", [])[:6]
    n = max(len(items), 1)
    cw = (SW - 1.0) / n
    cy = 2.35
    hline(slide, 0.5 + cw / 2, cy, cw * (n - 1), 1.5, LIGHT)
    for i, it in enumerate(items):
        cx = 0.5 + cw * i + cw / 2
        c = col(i)
        circle(slide, cx, cy, 0.9, fill=c, lbl=it.get("label", str(i + 1)), size=22)
        tag = shape(slide, MSO_SHAPE.PENTAGON, cx - 0.85, 3.05, 1.7, 0.45, fill=c)
        tag.rotation = 90 if False else 0
        label_in(tag, it.get("head", ""), 10)
        text(slide, cx - cw / 2 + 0.15, 3.7, cw - 0.3, 3.0, it.get("body", ""), 9, BODY, GRAY,
             align=PP_ALIGN.CENTER, line=1.3)


def lay_target(prs, slide, s):
    header(slide, s)
    cx, cy = 10.4, 4.15
    for i, d in enumerate([4.4, 3.5, 2.6, 1.7, 0.8]):
        circle(slide, cx, cy, d, fill=(RED if i % 2 == 0 else WHITE))
    items = s.get("items", [])[:4]
    y = 1.45
    for i, it in enumerate(items):
        c = col(i)
        circle(slide, 0.95, y + 0.42, 0.7, fill=c, lbl=it.get("icon", str(i + 1)), size=16)
        text(slide, 1.55, y + 0.08, 5.7, 0.3, it.get("head", ""), 12, XBOLD, INK)
        text(slide, 1.55, y + 0.42, 5.7, 0.9, it.get("body", ""), 9, BODY, GRAY, line=1.25)
        hline(slide, 1.55, y + 1.28, 5.7, 1.0, LIGHT)
        y += 1.45


def lay_numlist(prs, slide, s):
    header(slide, s)
    items = s.get("items", [])[:5]
    y = 1.5
    rh = min(1.15, 5.6 / max(len(items), 1))
    for i, it in enumerate(items):
        c = col(i)
        text(slide, 0.6, y + 0.1, 0.5, 0.4, f"{i+1:02d}", 14, XBOLD, c)
        text(slide, 1.3, y + 0.02, 4.2, 0.6, it.get("head", ""), 11.5, XBOLD, INK)
        text(slide, 5.7, y + 0.06, 6.2, rh - 0.25, it.get("body", ""), 9, BODY, GRAY, line=1.25)
        circle(slide, 12.45, y + rh / 2 - 0.08, 0.55, fill=c, lbl=it.get("icon", "●"), size=13)
        hline(slide, 0.6, y + rh - 0.12, 12.2, 1.0, LIGHT)
        y += rh


def lay_cards(prs, slide, s):
    header(slide, s)
    items = s.get("items", [])[:4]
    n = max(len(items), 1)
    gap = 0.3
    cw = (SW - 1.0 - gap * (n - 1)) / n
    for i, it in enumerate(items):
        x = 0.5 + i * (cw + gap)
        c = col(i)
        hd = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.5, cw, 0.75, fill=c)
        hd.adjustments[0] = 0.25
        label_in(hd, it.get("head", ""), 12)
        bx = shape(slide, MSO_SHAPE.RECTANGLE, x, 2.35, cw, 4.3, fill=LIGHT)
        text(slide, x + 0.2, 2.6, cw - 0.4, 3.8, it.get("body", ""), 9, BODY,
             RGBColor(0x5A, 0x5A, 0x66), line=1.35)


def lay_hub(prs, slide, s):
    header(slide, s)
    items = s.get("items", [])[:6]
    n = max(len(items), 1)
    cx, cy, R = SW / 2, 4.3, 1.85
    circle(slide, cx, cy, 1.5, fill=NAVY, lbl=s.get("center", str(n)), size=17)
    for i, it in enumerate(items):
        ang = -90 + i * (360 / n)
        rad = math.radians(ang)
        px, py = cx + R * math.cos(rad), cy + R * math.sin(rad)
        c = col(i)
        hline_x = px
        circle(slide, px, py, 0.6, fill=WHITE, line_color=c, line_w=2.0,
               lbl=f"{i+1:02d}", size=12, color=c)
        right = math.cos(rad) >= -0.15
        tw = 2.6
        tx = px + 0.55 if right else px - 0.55 - tw
        al = PP_ALIGN.LEFT if right else PP_ALIGN.RIGHT
        ty = py - 0.3
        text(slide, tx, ty, tw, 0.3, it.get("head", ""), 10.5, XBOLD, c, align=al)
        text(slide, tx, ty + 0.28, tw, 0.9, it.get("body", ""), 8.5, BODY, GRAY, align=al, line=1.2)


def lay_pyramid(prs, slide, s):
    header(slide, s)
    items = s.get("items", [])[:5]
    n = max(len(items), 1)
    base_w, top_w = 5.2, 1.6
    y0, rh, gap = 1.7, 0.85, 0.12
    cx = 4.0
    for i, it in enumerate(items):
        w = top_w + (base_w - top_w) * (i / max(n - 1, 1))
        c = col(i)
        bar = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, y0 + i * (rh + gap), w, rh, fill=c)
        bar.adjustments[0] = 0.5
        label_in(bar, it.get("head", f"Step {i+1:02d}"), 11)
        ty = y0 + i * (rh + gap) + 0.12
        circle(slide, 7.6, ty + 0.22, 0.45, fill=WHITE, line_color=c, line_w=1.75,
               lbl=it.get("icon", str(i + 1)), size=10)
        text(slide, 8.05, ty, 4.7, rh, it.get("body", ""), 9, BODY, GRAY, line=1.2)


def lay_stats(prs, slide, s):
    header(slide, s)
    items = s.get("items", [])[:4]
    n = max(len(items), 1)
    gap = 0.3
    cw = (SW - 1.0 - gap * (n - 1)) / n
    for i, it in enumerate(items):
        x = 0.5 + i * (cw + gap)
        c = col(i)
        shape(slide, MSO_SHAPE.RECTANGLE, x, 1.5, cw, 5.15, fill=LIGHT)
        text(slide, x, 1.85, cw, 0.8, it.get("value", ""), 34, XBOLD, c, align=PP_ALIGN.CENTER)
        text(slide, x, 2.7, cw, 0.3, it.get("head", ""), 10, SEMI, INK, align=PP_ALIGN.CENTER)
        pct = it.get("pct")
        pcy = 4.15
        if pct is not None:
            circle(slide, x + cw / 2, pcy, 1.3, fill=RGBColor(0xDD, 0xDD, 0xE3))
            pie_wedge(slide, x + cw / 2, pcy, 0.65, float(pct), c)
            circle(slide, x + cw / 2, pcy, 0.55, fill=LIGHT,
                   lbl=f"{pct}%", size=11, color=INK)
        text(slide, x + 0.2, 5.15, cw - 0.4, 1.3, it.get("body", ""), 8.5, BODY, GRAY,
             align=PP_ALIGN.CENTER, line=1.25)


def lay_gantt(prs, slide, s):
    header(slide, s)
    rows = s.get("rows", [])[:8]
    months = s.get("months", ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12"])
    x0, xg, w = 0.5, 3.6, SW - 0.5
    nm = len(months)
    cw = (w - xg) / nm
    text(slide, x0, 1.35, 3.0, 0.3, s.get("col_head", "TASK"), 9, XBOLD, GRAY, spacing=150)
    for j, m in enumerate(months):
        text(slide, xg + j * cw, 1.35, cw, 0.3, str(m), 9, SEMI, GRAY, align=PP_ALIGN.CENTER)
    y = 1.75
    rh = min(0.62, 5.3 / max(len(rows), 1))
    for i, r in enumerate(rows):
        c = col(i)
        if i % 2 == 0:
            shape(slide, MSO_SHAPE.RECTANGLE, x0, y, w - x0, rh, fill=LIGHT)
        text(slide, x0 + 0.1, y + rh / 2 - 0.14, 3.0, 0.3, r.get("task", ""), 9.5, SEMI, INK)
        a, b = int(r.get("start", 1)), int(r.get("end", 1))
        a = max(1, min(a, nm)); b = max(a, min(b, nm))
        bar = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, xg + (a - 1) * cw + 0.05,
                    y + rh / 2 - 0.11, (b - a + 1) * cw - 0.1, 0.22, fill=c)
        bar.adjustments[0] = 0.5
        y += rh
    for j in range(nm + 1):
        ln = slide.shapes.add_connector(1, Inches(xg + j * cw), Inches(1.7),
                                        Inches(xg + j * cw), Inches(y))
        ln.line.color.rgb = RGBColor(0xE4, 0xE4, 0xEA)
        ln.line.width = Pt(0.75)


def lay_matrix(prs, slide, s):
    header(slide, s)
    q = s.get("quads", [])[:4]
    cx, cy = SW / 2, 4.3
    qw, qh, gap = 3.6, 2.15, 0.15
    pos = [(cx - qw - gap / 2, cy - qh - gap / 2), (cx + gap / 2, cy - qh - gap / 2),
           (cx - qw - gap / 2, cy + gap / 2), (cx + gap / 2, cy + gap / 2)]
    for i, it in enumerate(q):
        x, y = pos[i]
        c = col(i)
        shape(slide, MSO_SHAPE.RECTANGLE, x, y, qw, qh, fill=LIGHT)
        shape(slide, MSO_SHAPE.RECTANGLE, x, y, qw, 0.08, fill=c)
        text(slide, x + 0.2, y + 0.22, qw - 0.4, 0.3, it.get("head", ""), 11, XBOLD, c)
        text(slide, x + 0.2, y + 0.58, qw - 0.4, qh - 0.7, it.get("body", ""), 8.5, BODY, GRAY, line=1.25)
    ax = s.get("axes", {})
    text(slide, cx - 2.0, cy - qh - gap / 2 - 0.35, 4.0, 0.25, ax.get("top", ""), 9, SEMI, GRAY, align=PP_ALIGN.CENTER)
    text(slide, cx - 2.0, cy + qh + gap / 2 + 0.08, 4.0, 0.25, ax.get("bottom", ""), 9, SEMI, GRAY, align=PP_ALIGN.CENTER)


def lay_table(prs, slide, s):
    header(slide, s)
    cols = s.get("cols", [])[:5]
    rows = s.get("rows", [])[:6]
    n = max(len(cols), 1)
    x0, w = 0.5, SW - 1.0
    cw = w / n
    y = 1.5
    for j, cname in enumerate(cols):
        hd = shape(slide, MSO_SHAPE.RECTANGLE, x0 + j * cw, y, cw - 0.06, 0.6, fill=col(j))
        label_in(hd, cname, 11)
    y += 0.68
    rh = min(0.85, 4.9 / max(len(rows), 1))
    for i, row in enumerate(rows):
        for j in range(n):
            cell = shape(slide, MSO_SHAPE.RECTANGLE, x0 + j * cw, y, cw - 0.06, rh - 0.08,
                         fill=(LIGHT if i % 2 == 0 else WHITE),
                         line_color=RGBColor(0xE4, 0xE4, 0xEA), line_w=0.75)
            val = row[j] if j < len(row) else ""
            tf = cell.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for m in ("margin_top", "margin_bottom"):
                setattr(tf, m, 0)
            tf.margin_left = Inches(0.12)
            tf.margin_right = Inches(0.12)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x5A, 0x5A, 0x66)
            _fonts(r, SEMI if j == 0 else BODY)
        y += rh


LAYOUTS = {
    "divider": lay_divider, "steps": lay_steps, "target": lay_target,
    "numlist": lay_numlist, "cards": lay_cards, "hub": lay_hub,
    "pyramid": lay_pyramid, "stats": lay_stats, "gantt": lay_gantt,
    "matrix": lay_matrix, "table": lay_table,
}


def build(spec_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation()
    prs.slide_width = Emu(int(SW * 914400))
    prs.slide_height = Emu(int(SH * 914400))
    blank = prs.slide_layouts[6]
    for i, s in enumerate(spec["slides"]):
        s["_page"] = i + 1
        slide = prs.slides.add_slide(blank)
        shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, fill=WHITE)
        fn = LAYOUTS.get(s.get("layout"))
        if not fn:
            raise SystemExit(f"unknown layout: {s.get('layout')}")
        fn(prs, slide, s)
    out = spec.get("output", "deck.pptx")
    prs.save(out)
    print(f"OK {out} ({len(spec['slides'])} slides)")


if __name__ == "__main__":
    build(sys.argv[1])
