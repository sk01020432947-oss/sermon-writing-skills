# -*- coding: utf-8 -*-
"""claude-ppt 렌더 엔진 — 검증된 최소 구현 (python-pptx 1.0+)

이 파일은 그대로 import 하거나 복사해 쓰는 참조 구현이다.
- 1000 x 563px 좌표계 (= 13.333 x 7.5in, 1px = 1/75in)
- 모든 도형의 테마 그림자 제거
- card_top_bar / card_left_bar : 액센트 3조각 규칙
- notes : 노트 마스터까지 폭을 맞춘 발표자 노트
한글은 Pretendard, 인용/거버닝은 NanumMyeongjo, 숫자/라틴은 Trebuchet MS.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from lxml import etree

# DATA = r"/Users/kbg1023/AppData\Local\Temp\claude\C--Users-ASUS--claude\25c091d2-bc9f-4391-beb9-557290afa5e4\scratchpad\series.json"
# OUT = r"/Users/kbg1023/Downloads\이교수의책과사람_에세이시리즈_87p_v2.pptx"

BG, BG_ALT = RGBColor(0xF0, 0xEE, 0xE6), RGBColor(0xF7, 0xF4, 0xEE)
CARD, CARD2 = RGBColor(0xE3, 0xDA, 0xCC), RGBColor(0xDD, 0xD6, 0xCE)
DARK = RGBColor(0x1C, 0x18, 0x10)
MAIN, MAIN_DEEP = RGBColor(0xCF, 0x6B, 0x4A), RGBColor(0xA5, 0x47, 0x2A)
MAIN_SOFT, MAIN_PALE = RGBColor(0xE0, 0x8B, 0x6A), RGBColor(0xED, 0xCB, 0xB8)
TINT = RGBColor(0xEA, 0xD6, 0xCC)
GRAY, RULE = RGBColor(0xB5, 0xAD, 0xA7), RGBColor(0xD0, 0xC9, 0xC2)
INK, INK2, INK3 = RGBColor(0x1C, 0x18, 0x10), RGBColor(0x4A, 0x45, 0x40), RGBColor(0x8A, 0x83, 0x7C)
POS = RGBColor(0x3D, 0x6B, 0x45)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DIM = RGBColor(0x5A, 0x54, 0x4E)

SANS, SERIF, NUM = "Pretendard", "NanumMyeongjo", "Trebuchet MS"
PXI = 1.0 / 75.0
def px(n): return Inches(n * PXI)
M, W = 45, 910

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def set_ea(run, face):
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(NS + tag)
        if el is None:
            el = etree.SubElement(rPr, NS + tag)
        el.set("typeface", face)


def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background(); r.shadow.inherit = False
    return s


def rect(s, x, y, w, h, fill=None, line=None, radius=None, lw=1, top_round=False):
    """top_round=True → 위 두 모서리만 둥근 면 (표 헤더·카드 상단 컬러바용)"""
    if top_round:
        geom = MSO_SHAPE.ROUND_2_SAME_RECTANGLE
    elif radius:
        geom = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        geom = MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(geom, px(x), px(y), px(w), px(h))
    if top_round:
        shp.adjustments[0] = min(0.5, (radius or 6) / min(w, h))
        shp.adjustments[1] = 0
    elif radius:
        shp.adjustments[0] = radius / min(w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    st = shp._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if st is not None:                      # 테마 effectRef(그림자)까지 제거
        shp._element.remove(st)
    return shp


def card_top_bar(s, x, y, w, h, accent, body=CARD, bar=7, r=6):
    """레퍼런스형 카드 — 상단 컬러 바(위 두 모서리 둥금) + 곧은 경계선 + 둥근 아래 모서리.
    바와 본문의 경계는 반드시 직선이어야 한다(카드를 둥근 채로 겹치면 옆구리에 색이 번진다)."""
    top = rect(s, x, y, w, 2 * r + bar, fill=accent, radius=r, top_round=True)
    rect(s, x, y + bar, w, h - bar - 2 * r, fill=body)              # 직선 경계 + 몸통
    bot = rect(s, x, y + h - 2 * r, w, 2 * r, fill=body, radius=r, top_round=True)
    bot.rotation = 180                                              # 아래 두 모서리만 둥글게
    return top


def _cap(s, x, y, w, h, fill, r, rot):
    """둥근 두 모서리만 가진 마감 조각. 90/270도는 회전 시 가로세로가 바뀌므로 미리 뒤집어 만든다."""
    if rot in (90, 270):
        cx, cy = x + w / 2, y + h / 2
        shp = rect(s, cx - h / 2, cy - w / 2, h, w, fill=fill, radius=r, top_round=True)
    else:
        shp = rect(s, x, y, w, h, fill=fill, radius=r, top_round=True)
    shp.rotation = rot
    return shp


def card_left_bar(s, x, y, w, h, accent, body=CARD, bar=4, r=5):
    """좌측 컬러 바 — 왼쪽 두 모서리만 둥글고, 바와 본문의 경계는 세로 직선."""
    _cap(s, x, y, bar + 2 * r, h, accent, r, 270)          # 왼쪽 둥근 마감 + 바
    rect(s, x + bar, y, w - bar - 2 * r, h, fill=body)      # 직선 경계 + 몸통
    _cap(s, x + w - 2 * r, y, 2 * r, h, body, r, 90)        # 오른쪽 둥근 마감


def text(s, x, y, w, h, runs, size=10.5, color=INK, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0, line=None):
    tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [(runs, {})]
    for i, item in enumerate(paras):
        t, o = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_before = Pt(o.get("space", space if i else 0))
        ls = o.get("line", line)
        if ls:
            p.line_spacing = ls
        r = p.add_run(); r.text = t
        f = r.font
        f.name = o.get("font", font); f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold); f.color.rgb = o.get("color", color)
        set_ea(r, o.get("font", font))
    return tb


PAGE = [0]
def header(s, crumb):
    text(s, M, 16, 520, 14, [(f"{crumb}   |   {PAGE[0]:02d}", {"color": INK2, "size": 8.5})])
    text(s, 1000 - M - 240, 16, 240, 14,
         [("이교수의 책과 사람", {"color": MAIN, "size": 8.5, "bold": True, "align": PP_ALIGN.RIGHT})])
    rect(s, M, 34, W, 0.8, fill=RULE)


def head(s, title, governing, tsize=27):
    text(s, M, 50, W, 40, title, size=tsize, bold=True, line=0.92)
    text(s, M, 88, W, 22, governing, size=11, color=INK2, font=SERIF, line=1.25)


def foot(s, note):
    text(s, M, 528, W, 14, note, size=7.5, color=INK3)


def fit(t, a=68, b=100):
    """길이에 따라 본문 크기를 낮춘다"""
    return 10.5 if len(t) <= a else (9.5 if len(t) <= b else 8.5)


NOTE_L, NOTE_W = Inches(0.75), Inches(6.0)   # 노트 페이지에서 슬라이드 축소본의 좌우 끝선

# 노트 페이지 크기(7.5 x 10in)와 노트 마스터 본문 상자를 먼저 못박는다.
# 마스터를 고치지 않으면 뷰에 따라 마스터 폭(더 넓음)을 따라가 노트가 슬라이드 밖으로 나간다.
_nsz = prs._element.find(
    "{http://schemas.openxmlformats.org/presentationml/2006/main}notesSz")
if _nsz is not None:
    _nsz.set("cx", str(Inches(7.5))); _nsz.set("cy", str(Inches(10)))
for _ph in prs.notes_master.placeholders:
    if _ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_IMAGE:
        _ph.left, _ph.top, _ph.width, _ph.height = NOTE_L, Inches(0.55), NOTE_W, Inches(3.375)
    elif _ph.placeholder_format.type == PP_PLACEHOLDER.BODY:
        _ph.left, _ph.top, _ph.width, _ph.height = NOTE_L, Inches(3.95), NOTE_W, Inches(5.55)


def notes(s, blocks, cap=1900):
    """발표자 노트 — 노트 페이지 좌우 여백까지 넓혀 본문을 최대한 담는다.
    blocks: [(소제목, 본문), ...]"""
    ns = s.notes_slide
    # 슬라이드 축소본과 노트 상자의 좌우 끝선을 같은 자리에 맞춘다 (노트 페이지 7.5 x 10in)
    for shp in ns.shapes:
        if shp.is_placeholder and shp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_IMAGE:
            shp.left, shp.top = NOTE_L, Inches(0.55)
            shp.width, shp.height = NOTE_W, Inches(3.375)  # 16:9 축소본 (6.0 x 3.375)
    ph = ns.notes_placeholder
    ph.left, ph.top = NOTE_L, Inches(3.95)
    ph.width, ph.height = NOTE_W, Inches(5.55)
    tf = ph.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    used, first = 0, True
    for htxt, body in blocks:
        if not body:
            continue
        body = body.strip()
        if used + len(body) > cap:
            body = body[: max(0, cap - used)].rstrip() + " …"
        items = ([(htxt, 9.5, True, RGBColor(0xA5, 0x47, 0x2A))] +
                 [(ln, 9.5, False, RGBColor(0x1C, 0x18, 0x10)) for ln in body.split("\n")])
        for txt, sz, bold, col in items:
            if not txt:
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(0 if bold and used == 0 else (8 if bold else 2))
            p.line_spacing = 1.22
            r = p.add_run(); r.text = txt
            r.font.size = Pt(sz); r.font.bold = bold
            r.font.color.rgb = col; r.font.name = SANS
            set_ea(r, SANS)
        used += len(body)
        if used >= cap:
            break


