# editorial-mono-deck builder: spec JSON -> editorial B/W pptx
# usage: python build.py spec.json
import json, os, sys, tempfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageOps

PAPER = RGBColor(0xF1, 0xEF, 0xE9)
INK = RGBColor(0x14, 0x14, 0x14)
GRAY = RGBColor(0x6E, 0x6C, 0x66)
PLACEHOLDER = RGBColor(0xD8, 0xD5, 0xCE)
DISPLAY = "Pretendard Black"
HEAVY = "Pretendard ExtraBold"
BODY = "Pretendard Regular"
MEDIUM = "Pretendard Medium"
SW, SH = 13.333, 7.5
_TMP = tempfile.mkdtemp(prefix="emd_")


def _fonts(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def _spc(run, val):  # tracking, 1/100 pt
    run._r.get_or_add_rPr().set("spc", str(val))


def text(slide, x, y, w, h, s, size=14, font=BODY, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, upper=False, spacing=None, line=1.0, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    lines = str(s).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run()
        r.text = ln.upper() if upper else ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)
        if spacing is not None:
            _spc(r, spacing)
    return tb


def rule(slide, x, y, w, weight=1.0, color=INK):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def rect(slide, x, y, w, h, fill=INK, line_color=None, line_w=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    return sp


def circle(slide, cx, cy, d, label, color=INK, filled=False, size=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    sp.shadow.inherit = False
    if filled:
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
    else:
        sp.fill.background()
    sp.line.color.rgb = color
    sp.line.width = Pt(1.75)
    tf = sp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(label)
    r.font.size = Pt(size or d * 26)
    r.font.color.rgb = (PAPER if filled else color)
    _fonts(r, HEAVY)
    return sp


def photo(slide, path, x, y, w, h, dark=False):
    """Grayscale center-crop fill; gray placeholder if path missing."""
    if path and os.path.isfile(path):
        im = Image.open(path)
        im = ImageOps.exif_transpose(im).convert("L")
        im = ImageOps.autocontrast(im, cutoff=1)
        tw, th = max(int(w * 150), 50), max(int(h * 150), 50)
        im = ImageOps.fit(im, (tw, th), Image.LANCZOS)
        out = os.path.join(_TMP, f"p{abs(hash((path, tw, th)))}.jpg")
        im.convert("RGB").save(out, quality=90)
        slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        rect(slide, x, y, w, h, fill=(RGBColor(0x2A, 0x2A, 0x28) if dark else PLACEHOLDER))


def meta_rows(slide, pairs, x, y, w, dark=False, row_h=0.42):
    fg = PAPER if dark else INK
    for i, (k, v) in enumerate(pairs):
        yy = y + i * row_h
        rule(slide, x, yy, w, 0.75, fg)
        text(slide, x, yy + 0.06, w * 0.42, 0.3, k, 9, MEDIUM, fg, upper=True, spacing=150)
        text(slide, x + w * 0.42, yy + 0.06, w * 0.58, 0.3, v, 9, BODY, fg, align=PP_ALIGN.RIGHT)
    rule(slide, x, y + len(pairs) * row_h, w, 0.75, fg)


def bg(slide, dark):
    rect(slide, 0, 0, SW, SH, fill=(INK if dark else PAPER))
    return (PAPER if dark else INK)


def pagenum(slide, spec, fg):
    n = spec.get("_page")
    if n:
        text(slide, SW - 1.0, SH - 0.45, 0.6, 0.3, f"{n:02d}", 9, MEDIUM, fg, align=PP_ALIGN.RIGHT, spacing=150)


# ---------------- layouts ----------------

def lay_cover(prs, slide, s):
    fg = bg(slide, s.get("dark", False))
    title = s.get("title", "TITLE")
    text(slide, 0.55, 0.4, 8.2, 4.6, title, s.get("title_size", 128), DISPLAY, fg, upper=True, spacing=-150, line=0.88)
    # vertical kicker bar: build horizontal, rotate 90 about center
    if s.get("kicker"):
        bw, bh, cx2, cy2 = 2.6, 0.55, 6.55, 3.55
        bar = rect(slide, cx2 - bw / 2, cy2 - bh / 2, bw, bh, fill=fg)
        tf = bar.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = s["kicker"].upper()
        r.font.size = Pt(13)
        r.font.color.rgb = PAPER if fg == INK else INK
        _fonts(r, HEAVY)
        _spc(r, 300)
        bar.rotation = 90
    photo(slide, s.get("image"), 8.75, 0.55, 4.05, 2.75, s.get("dark", False))
    if s.get("sub"):
        text(slide, 0.55, 5.35, 6.6, 1.4, s["sub"], 11.5, BODY, fg, line=1.3)
    meta_rows(slide, s.get("meta", []), 8.75, 3.85, 4.05, s.get("dark", False))
    if s.get("credit"):
        text(slide, 0.55, SH - 0.5, 6, 0.3, s["credit"], 9, MEDIUM, fg, upper=True, spacing=200)


def lay_toc(prs, slide, s):
    fg = bg(slide, s.get("dark", False))
    items = s.get("items", [])
    n = max(len(items), 1)
    cw = (SW - 1.1) / min(n, 6)
    for i, it in enumerate(items[:6]):
        x = 0.55 + i * cw
        text(slide, x, 0.5, cw - 0.25, 0.3, f"{i+1:02d}", 10, HEAVY, fg, spacing=150)
        text(slide, x, 0.82, cw - 0.25, 0.55, it, 10, MEDIUM, fg, upper=True, spacing=100)
        rule(slide, x, 1.42, cw - 0.25, 0.75, fg)
    text(slide, 0.55, 2.2, SW - 1.1, 1.9, s.get("title", "CONTENTS"), 110, DISPLAY, fg,
         align=PP_ALIGN.CENTER, upper=True, spacing=-100)
    photo(slide, s.get("image"), (SW - 4.6) / 2, 4.35, 4.6, 2.65, s.get("dark", False))
    pagenum(slide, s, fg)


def lay_section(prs, slide, s):
    dark = s.get("dark", True)
    fg = bg(slide, dark)
    photo(slide, s.get("image"), 0.55, 0.55, 3.5, 3.5, dark)
    if s.get("body"):
        text(slide, 4.55, 0.8, 5.6, 3.2, s["body"], 11.5, BODY, fg, line=1.35)
    meta_rows(slide, s.get("meta", []), 10.6, 0.8, 2.2, dark)
    text(slide, 0.55, 4.55, 9.6, 2.4, s.get("title", "SECTION"), 100, DISPLAY, fg, upper=True,
         anchor=MSO_ANCHOR.BOTTOM, spacing=-100)
    if s.get("number"):
        text(slide, 10.4, 4.55, 2.4, 2.4, str(s["number"]), 100, DISPLAY, fg,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)
    pagenum(slide, s, fg)


def lay_columns(prs, slide, s):
    fg = bg(slide, s.get("dark", False))
    text(slide, 0.55, 0.45, 6.5, 1.1, s.get("title", ""), 60, DISPLAY, fg, upper=True, spacing=-80)
    if s.get("note"):
        rect(slide, 8.9, 0.45, 3.9, 1.35, fill=None, line_color=fg, line_w=1.0)
        text(slide, 9.1, 0.62, 3.5, 1.0, s["note"], 8.5, BODY, fg, line=1.25)
    photo(slide, s.get("image"), 0.55, 1.9, 2.9, 2.5, s.get("dark", False))
    if s.get("summary"):
        text(slide, 0.55, 4.6, 2.9, 2.4, s["summary"], 9.5, BODY, fg, line=1.3)
    items = s.get("items", [])[:4]
    n = max(len(items), 1)
    x0, tw = 3.85, SW - 3.85 - 0.55
    cw = tw / n
    for i, it in enumerate(items):
        x = x0 + i * cw
        circle(slide, x + 0.35, 2.25, 0.7, it.get("label", str(i + 1)), fg, size=20)
        text(slide, x + 0.85, 2.0, cw - 1.05, 0.35, it.get("head", ""), 12, HEAVY, fg, upper=True)
        text(slide, x + 0.85, 2.32, cw - 1.05, 0.3, it.get("sub", ""), 8.5, MEDIUM, GRAY if fg == INK else PAPER, upper=True, spacing=100)
        rule(slide, x, 2.85, cw - 0.3, 0.75, fg)
        text(slide, x, 3.05, cw - 0.3, 3.8, it.get("body", ""), 9.5, BODY, fg, line=1.3)
    pagenum(slide, s, fg)


def lay_feature(prs, slide, s):
    dark = s.get("dark", False)
    fg = bg(slide, dark)
    photo(slide, s.get("image"), 0.55, 0.9, 3.3, 4.3, dark)
    if s.get("intro"):
        text(slide, 0.55, 0.35, 7.5, 0.5, s["intro"], 8.5, BODY, fg, line=1.2)
    x0 = 4.35
    w = SW - x0 - 0.55
    if s.get("badge"):
        circle(slide, x0 + 0.55, 1.45, 1.1, s["badge"], fg, size=30)
    hx = x0 + 1.5
    rule(slide, hx, 1.05, w - 1.5, 1.25, fg)
    text(slide, hx, 1.2, w - 3.4, 0.55, s.get("head", ""), 26, DISPLAY, fg, upper=True, spacing=-50)
    text(slide, hx + (w - 3.3), 1.28, 1.8, 0.4, s.get("period", ""), 15, HEAVY, fg, align=PP_ALIGN.RIGHT)
    rule(slide, hx, 1.85, w - 1.5, 1.25, fg)
    if s.get("sub"):
        text(slide, hx, 2.0, w - 1.5, 0.35, s["sub"], 11, HEAVY, fg)
        rule(slide, hx, 2.42, w - 1.5, 0.75, fg)
    cols = s.get("cols", [])
    ncol = max(len(cols), 1)
    cw = (w - 0.3 * (ncol - 1)) / ncol if ncol else w
    for i, c in enumerate(cols[:3]):
        text(slide, x0 + i * (cw + 0.3), 2.6, cw, 2.5, c, 9.5, BODY, fg, line=1.3)
    text(slide, 0.55, 5.15, SW - 1.1, 2.0, s.get("title", ""), 84, DISPLAY, fg, upper=True,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM, spacing=-100)
    pagenum(slide, s, fg)


def lay_listphoto(prs, slide, s):
    fg = bg(slide, s.get("dark", False))
    if s.get("intro"):
        text(slide, 0.55, 0.4, 5.5, 0.9, s["intro"], 9, BODY, fg, line=1.25)
    photo(slide, s.get("image"), 2.8, 1.15, 3.6, 3.3, s.get("dark", False))
    text(slide, 0.55, 4.7, 7.2, 2.1, s.get("title", ""), 88, DISPLAY, fg, upper=True,
         anchor=MSO_ANCHOR.BOTTOM, spacing=-100)
    items = s.get("items", [])[:4]
    x0, w = 7.0, SW - 7.0 - 0.55
    y = 0.7
    for i, it in enumerate(items):
        circle(slide, x0 + 0.3, y + 0.3, 0.6, it.get("label", str(i + 1)), fg, size=17)
        text(slide, x0 + 0.85, y + 0.06, 1.9, 0.5, it.get("head", ""), 11, HEAVY, fg, upper=True)
        text(slide, x0 + 2.85, y, w - 2.85, 1.35, it.get("body", ""), 9, BODY, fg, line=1.25)
        y += max(1.0, 0.35 + 0.22 * (len(it.get("body", "")) // 46 + 1))
        rule(slide, x0, y - 0.18, w, 0.75, fg)
    pagenum(slide, s, fg)


def lay_table(prs, slide, s):
    fg = bg(slide, s.get("dark", False))
    if s.get("intro"):
        text(slide, 0.55, 0.4, 4.6, 0.8, s["intro"], 9, BODY, fg, line=1.25)
    rows = s.get("rows", [])[:6]
    x0, w, y = 0.55, 4.9, 1.5
    rule(slide, x0, y, w, 1.25, fg)
    for name, val in rows:
        text(slide, x0, y + 0.12, w * 0.7, 0.35, name, 13, HEAVY, fg, upper=True, spacing=50)
        text(slide, x0 + w * 0.7, y + 0.12, w * 0.3, 0.35, val, 13, HEAVY, fg, align=PP_ALIGN.RIGHT)
        y += 0.52
        rule(slide, x0, y, w, 0.75, fg)
    if s.get("footnote"):
        text(slide, 0.55, y + 0.3, 4.6, 1.2, s["footnote"], 8.5, BODY, GRAY if fg == INK else PAPER, line=1.25)
    photo(slide, s.get("image"), 6.3, 0.55, 4.1, 4.3, s.get("dark", False))
    parts = (s.get("title", "") or "").split(" ", 1)
    text(slide, 0.55, 5.2, 6.0, 2.0, parts[0], 80, DISPLAY, fg, upper=True, anchor=MSO_ANCHOR.BOTTOM,
         spacing=-100, wrap=False)
    if len(parts) > 1:
        text(slide, 6.6, 5.2, 6.2, 2.0, parts[1], 80, DISPLAY, fg, upper=True,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM, spacing=-100, wrap=False)
    pagenum(slide, s, fg)


def lay_strip(prs, slide, s):
    fg = bg(slide, s.get("dark", False))
    parts = (s.get("title", "") or "").split(" ", 1)
    text(slide, 0.55, 0.4, 6.5, 1.3, parts[0], 66, DISPLAY, fg, upper=True, spacing=-80)
    if len(parts) > 1:
        text(slide, 6.3, 0.4, 6.5, 1.3, parts[1], 66, DISPLAY, fg, upper=True, align=PP_ALIGN.RIGHT, spacing=-80)
    if s.get("intro"):
        text(slide, 0.55, 1.85, 6.8, 0.9, s["intro"], 9.5, BODY, fg, line=1.3)
    cells = s.get("cells", [])[:4]
    n = max(len(cells), 1)
    cw = (SW - 1.1 - 0.35 * (n - 1)) / n
    for i, c in enumerate(cells):
        x = 0.55 + i * (cw + 0.35)
        text(slide, x, 3.0, cw, 0.3, c.get("head", ""), 10, HEAVY, fg, upper=True, spacing=100)
        rule(slide, x, 3.35, cw, 0.75, fg)
        photo(slide, c.get("image"), x, 3.55, cw, 2.55, s.get("dark", False))
        if c.get("body"):
            text(slide, x, 6.2, cw, 0.9, c["body"], 8.5, BODY, fg, line=1.2)
    pagenum(slide, s, fg)


def lay_closing(prs, slide, s):
    fg = bg(slide, s.get("dark", False))
    text(slide, 0.55, 0.45, 10.2, 1.7, s.get("title", "GET IN TOUCH"), 88, DISPLAY, fg, upper=True, spacing=-100)
    if s.get("badge"):
        circle(slide, 11.9, 1.3, 1.15, s["badge"], fg, size=32)
    rule(slide, 0.55, 2.35, SW - 1.1, 1.25, fg)
    meta_rows(slide, s.get("meta", []), 0.55, 2.9, 5.4, s.get("dark", False), row_h=0.55)
    if s.get("body"):
        text(slide, 0.55, 5.6, 5.4, 1.5, s["body"], 9.5, BODY, fg, line=1.3)
    photo(slide, s.get("image"), 7.0, 2.85, 5.75, 4.1, s.get("dark", False))
    pagenum(slide, s, fg)


LAYOUTS = {
    "cover": lay_cover, "toc": lay_toc, "section": lay_section, "columns": lay_columns,
    "feature": lay_feature, "listphoto": lay_listphoto, "table": lay_table,
    "strip": lay_strip, "closing": lay_closing,
}


def build(spec_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation()
    prs.slide_width = Emu(int(SW * 914400))
    prs.slide_height = Emu(int(SH * 914400))
    blank = prs.slide_layouts[6]
    for i, s in enumerate(spec["slides"]):
        s["_page"] = i + 1
        slide = prs.slides.add_slide(blank)
        fn = LAYOUTS.get(s.get("layout"))
        if not fn:
            raise SystemExit(f"unknown layout: {s.get('layout')}")
        fn(prs, slide, s)
    out = spec.get("output", "deck.pptx")
    prs.save(out)
    print(f"OK {out} ({len(spec['slides'])} slides)")


if __name__ == "__main__":
    build(sys.argv[1])
