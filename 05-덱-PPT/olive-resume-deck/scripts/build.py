# olive-resume-deck builder: spec JSON -> cream/olive/ink retro resume pptx
# usage: python build.py spec.json
import json, os, sys, tempfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image, ImageOps

CREAM = RGBColor(0xEF, 0xE8, 0xDB)
INK = RGBColor(0x16, 0x16, 0x16)
OLIVE = RGBColor(0xA9, 0xA2, 0x57)
DARKP = RGBColor(0x1C, 0x1C, 0x1C)
GRAY = RGBColor(0x6E, 0x6A, 0x60)
PLACEHOLDER = RGBColor(0xD8, 0xD2, 0xC2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK_F = "Pretendard Black"
XBOLD = "Pretendard ExtraBold"
SEMI = "Pretendard SemiBold"
BODY = "Pretendard Regular"
SW, SH = 13.333, 7.5
_TMP = tempfile.mkdtemp(prefix="ord_")


def _fonts(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def text(slide, x, y, w, h, s, size=9, font=BODY, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line=1.25, wrap=True, spacing=None, upper=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        r = p.add_run()
        r.text = ln.upper() if upper else ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)
        if spacing is not None:
            r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def rect(slide, x, y, w, h, fill=INK, line_color=None, line_w=1.0, round_=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    return sp


def label_in(sp, s, size=9, color=WHITE, font=XBOLD):
    tf = sp.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        _fonts(r, font)


def circle(slide, cx, cy, d, fill=INK, lbl=None, size=None, color=WHITE, line_color=None, line_w=1.5):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    if lbl is not None:
        label_in(sp, lbl, size or max(9, d * 16), color)
    return sp


def hline(slide, x, y, w, weight=1.0, color=INK):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def photo(slide, path, x, y, w, h, dark=False):
    if path and os.path.isfile(path):
        im = Image.open(path)
        im = ImageOps.exif_transpose(im).convert("RGB")
        tw, th = max(int(w * 150), 50), max(int(h * 150), 50)
        im = ImageOps.fit(im, (tw, th), Image.LANCZOS)
        out = os.path.join(_TMP, f"p{abs(hash((path, tw, th)))}.jpg")
        im.save(out, quality=90)
        slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        rect(slide, x, y, w, h, fill=(RGBColor(0x33, 0x33, 0x30) if dark else PLACEHOLDER))


def pills(slide, tags, x, y, max_w, dark=False, filled_idx=0):
    """Row of pill chips; wraps to next row when exceeding max_w."""
    cx, cy = x, y
    for i, t in enumerate(tags):
        w = 0.28 + len(str(t)) * 0.085
        if cx + w > x + max_w:
            cx = x
            cy += 0.42
        filled = (i % 3 == filled_idx % 3)
        sp = rect(slide, cx, cy, w, 0.3, fill=(OLIVE if filled else None),
                  line_color=(None if filled else (WHITE if dark else INK)),
                  line_w=0.75, round_=True)
        sp.adjustments[0] = 0.5
        label_in(sp, t, 7.5, (INK if filled else (WHITE if dark else INK)), SEMI)
        cx += w + 0.14
    return cy + 0.42


def chrome(slide, spec, s, dark=False, side_color=None, side_y=1.2, side_h=4.5):
    """Footer + right vertical text on every slide."""
    fg = WHITE if dark else INK
    hline(slide, 0.45, SH - 0.42, SW - 1.2, 0.75, fg)
    text(slide, 0.45, SH - 0.34, 4.0, 0.25, spec.get("address", ""), 6.5, BODY, fg)
    text(slide, SW / 2 - 2.0, SH - 0.34, 4.0, 0.25, "DESIGN BY " + spec.get("author", "").upper(),
         6.5, SEMI, fg, align=PP_ALIGN.CENTER)
    n = s.get("_page")
    if n:
        text(slide, SW - 1.15, SH - 0.34, 0.4, 0.25, f"{n:02d}", 7, XBOLD, fg, align=PP_ALIGN.RIGHT)
    side = spec.get("side", "")
    if side:
        fg = side_color or fg
        vt = rect(slide, SW - 0.52, side_y, 0.4, side_h, fill=None)
        tf = vt.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = side
        r.font.size = Pt(7)
        r.font.color.rgb = fg
        _fonts(r, SEMI)
        r._r.get_or_add_rPr().set("spc", "300")
        tf._txBody.find(qn("a:bodyPr")).set("vert", "vert270")


# ---------------- layouts ----------------

def lay_cover(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    rect(slide, 0, 0, SW, 4.9, fill=DARKP)
    photo(slide, s.get("image"), 7.6, 0.35, 4.6, 4.2)
    text(slide, 0.5, 0.25, 7.4, 2.3, s.get("title", "HASU"), 120, BLACK_F, WHITE, upper=True,
         spacing=-150, line=0.9, wrap=False)
    if s.get("mark"):
        text(slide, 6.55, 0.5, 0.8, 0.5, s["mark"], 20, BLACK_F, WHITE)
    ab = rect(slide, 0.5, 2.85, 3.4, 0.75, fill=OLIVE)
    label_in(ab, s.get("accent", "RESUME"), 30, INK, BLACK_F)
    if s.get("sub"):
        text(slide, 0.5, 3.75, 5.5, 0.35, s["sub"], 11, SEMI, WHITE, upper=True, spacing=250)
    for i, bw in enumerate([0.04, 0.02, 0.05, 0.02, 0.03, 0.05, 0.02, 0.04, 0.03, 0.05, 0.02, 0.03]):
        rect(slide, 4.6 + i * 0.1, 2.95, bw, 0.4, fill=WHITE)
    if s.get("body"):
        text(slide, 0.5, 5.35, 6.5, 1.2, s["body"], 9, BODY, INK, line=1.35)
    chrome(slide, spec, s, side_color=WHITE, side_y=0.6, side_h=3.9)


def lay_toc(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    text(slide, 0.5, 0.5, 7.5, 1.0, s.get("title", "CONTENTS LIST"), 40, BLACK_F, INK, upper=True, spacing=-50)
    photo(slide, s.get("image"), 7.3, 0.6, 5.0, 3.6)
    items = s.get("items", [])[:7]
    y = 1.8
    for i, it in enumerate(items):
        head = it.get("head", it) if isinstance(it, dict) else it
        body = it.get("body", "") if isinstance(it, dict) else ""
        text(slide, 0.5, y, 0.7, 0.3, f"{i+1:02d}.", 10, BLACK_F, INK)
        text(slide, 1.3, y, 3.4, 0.3, head, 11, XBOLD, INK, upper=True)
        if body:
            text(slide, 1.3, y + 0.3, 5.4, 0.3, body, 7.5, BODY, GRAY)
        text(slide, 6.3, y, 0.5, 0.3, f"{i+2:02d}", 9, SEMI, GRAY, align=PP_ALIGN.RIGHT)
        hline(slide, 0.5, y + 0.62, 6.3, 0.75, INK)
        y += 0.72
    chrome(slide, spec, s)


def lay_about(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    photo(slide, s.get("image"), 0.5, 0.55, 4.4, 5.5)
    if s.get("body"):
        text(slide, 5.5, 0.7, 6.9, 2.2, s["body"], 9, BODY, INK, line=1.4)
    if s.get("tags"):
        pills(slide, s["tags"][:6], 5.5, 3.1, 6.5)
    text(slide, 5.5, 4.2, 7.0, 2.0, s.get("title", "ABOUT\nME"), 54, BLACK_F, INK, upper=True,
         spacing=-80, line=0.95)
    chrome(slide, spec, s)


def lay_columns(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    text(slide, 0.5, 0.4, 8.0, 1.1, s.get("title", ""), 46, BLACK_F, INK, upper=True, spacing=-60)
    if s.get("body"):
        text(slide, 0.5, 1.5, 7.6, 0.7, s["body"], 8.5, BODY, GRAY, line=1.3)
    photo(slide, s.get("image"), 9.0, 0.5, 3.3, 3.4)
    items = s.get("items", [])[:4]
    n = max(len(items), 1)
    x0, tw = 0.5, 8.2
    cw = tw / n
    y0 = 2.5
    for i, it in enumerate(items):
        x = x0 + i * cw
        text(slide, x, y0, cw - 0.25, 0.3, it.get("head", ""), 10.5, BLACK_F, INK, upper=True)
        hline(slide, x, y0 + 0.35, cw - 0.25, 1.25, INK)
        text(slide, x, y0 + 0.5, cw - 0.25, 2.6, it.get("body", ""), 7.5, BODY, GRAY, line=1.3)
        if it.get("tag"):
            tg = rect(slide, x, y0 + 3.35, min(cw - 0.35, 1.15), 0.3, fill=OLIVE, round_=True)
            tg.adjustments[0] = 0.5
            label_in(tg, it["tag"], 7.5, INK, SEMI)
    chrome(slide, spec, s)


def lay_card(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    rect(slide, 0.45, 0.45, SW - 0.95, SH - 1.15, fill=OLIVE)
    text(slide, 0.8, 0.7, 8.0, 1.0, s.get("title", ""), 44, BLACK_F, INK, upper=True, spacing=-60)
    photo(slide, s.get("image"), 0.8, 2.0, 3.2, 4.3)
    hx, hw = 4.35, SW - 4.35 - 0.85
    hline(slide, hx, 2.1, hw, 1.25, INK)
    text(slide, hx, 2.25, hw - 2.0, 0.4, s.get("head", ""), 19, BLACK_F, INK, upper=True)
    text(slide, hx + hw - 1.9, 2.3, 1.9, 0.35, s.get("period", ""), 14, XBOLD, INK, align=PP_ALIGN.RIGHT)
    hline(slide, hx, 2.75, hw, 1.25, INK)
    if s.get("sub"):
        text(slide, hx, 2.9, hw, 0.3, s["sub"], 10, SEMI, INK, upper=True, spacing=100)
        hline(slide, hx, 3.25, hw, 0.75, INK)
    if s.get("body"):
        text(slide, hx, 3.5, hw * 0.52, 2.9, s["body"], 8.5, BODY, INK, line=1.35)
    bullets = s.get("bullets", [])[:6]
    by = 3.5
    for b in bullets:
        text(slide, hx + hw * 0.56, by, hw * 0.44, 0.35, "•  " + b, 8.5, SEMI, INK, line=1.2)
        by += 0.42
    chrome(slide, spec, s)


def lay_circles(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    items = s.get("items", [])[:4]
    n = max(len(items), 1)
    x0, tw = 0.5, 8.4
    cw = tw / n
    for i, it in enumerate(items):
        x = x0 + i * cw
        c = INK if i % 2 == 0 else OLIVE
        tc = WHITE if i % 2 == 0 else INK
        circle(slide, x + 0.75, 1.15, 1.3, fill=c, lbl=it.get("label", f"{i+1:02d}"), size=13, color=tc)
        text(slide, x, 2.15, cw - 0.25, 0.3, it.get("head", ""), 9.5, BLACK_F, INK, upper=True)
        hline(slide, x, 2.5, cw - 0.25, 1.0, INK)
        text(slide, x, 2.65, cw - 0.25, 2.2, it.get("body", ""), 7.5, BODY, GRAY, line=1.3)
        if it.get("tag"):
            tg = rect(slide, x, 4.9, min(cw - 0.35, 1.15), 0.3, fill=OLIVE, round_=True)
            tg.adjustments[0] = 0.5
            label_in(tg, it["tag"], 7.5, INK, SEMI)
    photo(slide, s.get("image"), 9.3, 0.55, 3.0, 3.9)
    if s.get("body"):
        text(slide, 0.5, 5.45, 7.0, 0.6, s["body"], 8, BODY, GRAY, line=1.3)
    text(slide, 7.6, 5.35, 5.2, 1.2, s.get("title", ""), 44, BLACK_F, INK, upper=True,
         align=PP_ALIGN.RIGHT, spacing=-60, anchor=MSO_ANCHOR.BOTTOM, wrap=False)
    chrome(slide, spec, s)


def lay_darktable(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    text(slide, 0.5, 0.45, 7.8, 1.9, s.get("title", ""), 44, BLACK_F, INK, upper=True,
         spacing=-60, line=0.95)
    photo(slide, s.get("image"), 0.5, 2.6, 3.4, 4.0)
    px, pw = 4.15, SW - 4.15 - 0.6
    rect(slide, px, 2.3, pw, 4.35, fill=DARKP)
    text(slide, px + 0.35, 2.6, pw - 0.7, 0.45, s.get("panel_head", ""), 17, BLACK_F, WHITE, upper=True)
    hline(slide, px + 0.35, 3.15, pw - 0.7, 1.0, OLIVE)
    groups = s.get("groups", [])[:2]
    gw = (pw - 1.05) / max(len(groups), 1)
    for gi, g in enumerate(groups):
        gx = px + 0.35 + gi * (gw + 0.35)
        text(slide, gx, 3.35, gw, 0.3, g.get("head", ""), 9.5, SEMI, OLIVE, upper=True, spacing=100)
        y = 3.8
        for name, val in g.get("rows", [])[:6]:
            text(slide, gx, y, gw * 0.7, 0.28, name, 8.5, SEMI, WHITE)
            text(slide, gx + gw * 0.7, y, gw * 0.3, 0.28, val, 8.5, XBOLD, OLIVE, align=PP_ALIGN.RIGHT)
            y += 0.42
    chrome(slide, spec, s)


def lay_stats(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    text(slide, 0.5, 0.45, 6.5, 1.9, s.get("title", ""), 44, BLACK_F, INK, upper=True,
         spacing=-60, line=0.95)
    if s.get("body"):
        text(slide, 7.3, 0.6, 5.3, 1.5, s["body"], 8.5, BODY, GRAY, line=1.35)
    photo(slide, s.get("image"), 0.5, 2.6, 3.1, 4.0)
    items = s.get("items", [])[:4]
    n = max(len(items), 1)
    x0, tw = 4.0, SW - 4.0 - 0.6
    cw = tw / n
    for i, it in enumerate(items):
        x = x0 + i * cw
        circle(slide, x + cw / 2, 3.35, 1.35, fill=INK, lbl=it.get("value", ""), size=16)
        text(slide, x + 0.1, 4.35, cw - 0.2, 1.5, it.get("body", ""), 7.5, BODY, GRAY,
             align=PP_ALIGN.CENTER, line=1.3)
        text(slide, x + 0.1, 6.0, cw - 0.2, 0.3, it.get("head", ""), 9, BLACK_F, INK,
             align=PP_ALIGN.CENTER, upper=True)
    chrome(slide, spec, s)


def lay_sidelist(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    if s.get("body"):
        text(slide, 0.5, 0.55, 6.6, 1.5, s["body"], 8.5, BODY, GRAY, line=1.35)
    if s.get("tags"):
        pills(slide, s["tags"][:6], 0.5, 2.25, 6.2)
    text(slide, 0.5, 3.3, 7.0, 2.2, s.get("title", ""), 50, BLACK_F, INK, upper=True,
         spacing=-70, line=0.95)
    items = s.get("items", [])[:3]
    y = 0.7
    for i, it in enumerate(items):
        photo(slide, it.get("image"), 8.0, y, 1.5, 1.75)
        text(slide, 9.7, y + 0.05, 2.9, 0.3, it.get("head", ""), 9.5, BLACK_F, INK, upper=True)
        text(slide, 9.7, y + 0.38, 2.9, 1.3, it.get("body", ""), 7.5, BODY, GRAY, line=1.3)
        y += 2.0
    chrome(slide, spec, s)


def lay_collage(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    text(slide, 0.5, 0.5, 6.0, 1.1, s.get("title", ""), 50, BLACK_F, INK, upper=True, spacing=-70)
    if s.get("body"):
        text(slide, 0.5, 1.8, 5.8, 2.6, s["body"], 8.5, BODY, GRAY, line=1.4)
    if s.get("tags"):
        pills(slide, s["tags"][:8], 0.5, 4.7, 5.6)
    imgs = s.get("images", [])
    cells = [(6.7, 0.5, 2.7, 3.1), (9.55, 0.5, 2.75, 2.0), (9.55, 2.65, 2.75, 2.0),
             (6.7, 3.75, 2.7, 2.9), (9.55, 4.8, 2.75, 1.85)]
    for i, (x, y, w, h) in enumerate(cells):
        photo(slide, imgs[i] if i < len(imgs) else None, x, y, w, h)
    chrome(slide, spec, s)


def lay_contact(prs, slide, spec, s):
    rect(slide, 0, 0, SW, SH, fill=CREAM)
    photo(slide, s.get("image"), 0.5, 0.55, 5.4, 5.9)
    text(slide, 6.4, 0.8, 6.3, 2.0, s.get("title", "GET\nIN TOUCH"), 50, BLACK_F, INK, upper=True,
         spacing=-70, line=0.95)
    cols = s.get("meta", [])[:3]
    cw = (SW - 6.4 - 0.7) / max(len(cols), 1)
    for i, (k, v) in enumerate(cols):
        x = 6.4 + i * cw
        text(slide, x, 3.2, cw - 0.2, 0.3, k, 9, BLACK_F, INK, upper=True, spacing=100)
        text(slide, x, 3.55, cw - 0.2, 1.0, v, 8, BODY, GRAY, line=1.3)
    if s.get("tags"):
        pills(slide, s["tags"][:6], 6.4, 4.9, 6.0)
    chrome(slide, spec, s)


LAYOUTS = {
    "cover": lay_cover, "toc": lay_toc, "about": lay_about, "columns": lay_columns,
    "card": lay_card, "circles": lay_circles, "darktable": lay_darktable, "stats": lay_stats,
    "sidelist": lay_sidelist, "collage": lay_collage, "contact": lay_contact,
}


def build(spec_path):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    prs = Presentation()
    prs.slide_width = Emu(int(SW * 914400))
    prs.slide_height = Emu(int(SH * 914400))
    blank = prs.slide_layouts[6]
    for i, s in enumerate(spec["slides"]):
        s["_page"] = i + 1
        slide = prs.slides.add_slide(blank)
        fn = LAYOUTS.get(s.get("layout"))
        if not fn:
            raise SystemExit(f"unknown layout: {s.get('layout')}")
        fn(prs, slide, spec, s)
    out = spec.get("output", "deck.pptx")
    prs.save(out)
    print(f"OK {out} ({len(spec['slides'])} slides)")


if __name__ == "__main__":
    build(sys.argv[1])
